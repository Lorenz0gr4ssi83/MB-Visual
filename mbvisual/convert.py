"""Convertitore universale: "Salva con nome" in quasi ogni formato.

Per ogni tipo di file indica i formati di destinazione possibili e sceglie il motore migliore:
- Word / Excel / PowerPoint installati (automazione COM: massima fedelta'), altrimenti LibreOffice
- Pillow (immagini), FFmpeg (audio e video), PyMuPDF (PDF, SVG, XPS, ePub, HTML), pdf2docx (PDF -> Word)
- 7-Zip (archivi), fontTools (font), sqlite3 + openpyxl (database), markdown (Markdown -> HTML)
"""
import csv
import html as htmlmod
import io
import json
import os
import shutil
import sqlite3
import tempfile
import zipfile
from pathlib import Path

from . import misc
from .paths import IS_WIN, NO_WINDOW, TEMP_DIR, find_libreoffice, find_office as _find_office
from .i18n import t as tr

# famiglie di formati di partenza (estensione -> famiglia)
FAMILIES = {
    "word": "doc docx docm dot dotx dotm rtf odt ott fodt wpd wps".split(),
    "sheet": "xls xlsx xlsm xlsb xlt xltx xltm ods fods ots csv tsv dif slk sylk prn".split(),
    "slides": "ppt pptx pptm pps ppsx ppsm pot potx potm odp otp fodp".split(),
    "pdf": ["pdf"],
    "xps": ["xps", "oxps", "cbz", "fb2", "mobi"],
    "epub": ["epub"],
    "svg": ["svg", "svgz"],
    "html": "html htm xhtml mht mhtml".split(),
    "markdown": "md markdown mdown mkd".split(),
    "image": list(misc.PIL_FORMATS) + sorted(misc.RAW_EXTS) + ["psd", "pcd", "fits", "fit", "xpm", "wmf", "emf", "cur"],
    "media": "mp4 m4v mov qt webm mkv avi wmv asf flv f4v mpg mpeg ts mts m2ts 3gp 3g2 ogv vob divx".split(),
    "audio": "mp3 wav wave flac ogg oga opus m4a m4b aac wma aif aiff aifc ac3 amr ape wv mka".split(),
    "archive": "zip 7z rar tar gz tgz bz2 tbz2 xz txz lzma cab iso wim arj lzh cpio rpm deb jar".split(),
    "font": "ttf otf woff woff2".split(),
    "database": "sqlite sqlite3 db db3 s3db sl3".split(),
    "text": "txt text log ini cfg conf json xml yaml yml toml csv tsv sql py js ts css c h cpp cs java bat ps1 sh".split(),
    "email": ["eml", "msg"],
}
FAMILY_OF = {}
for fam, exts in FAMILIES.items():
    for e in exts:
        FAMILY_OF.setdefault(e, fam)

# etichette: chiavi di traduzione (ui/i18n), tradotte in targets() nella lingua della richiesta
TARGETS = {
    "word": [("docx", "Word"), ("doc", "srv_convert.word_97_2003"), ("docm", "srv_convert.word_macro"), ("dotx", "srv_convert.modello_word"), ("pdf", "PDF"),
             ("odt", "srv_convert.opendocument_testo"), ("rtf", "RTF"), ("txt", "srv_convert.testo_semplice"), ("html", "srv_convert.pagina_web"), ("mht", "srv_convert.pagina_web_file"),
             ("xps", "XPS"), ("xml", "srv_convert.word_xml"), ("epub", "eBook")],
    "sheet": [("xlsx", "Excel"), ("xls", "srv_convert.excel_97_2003"), ("xlsm", "srv_convert.excel_macro"), ("xlsb", "srv_convert.excel_binario"), ("xltx", "srv_convert.modello_excel"),
              ("ods", "srv_convert.opendocument_foglio"), ("csv", "CSV (UTF-8)"), ("txt", "srv_convert.testo_separato_tabulazioni"), ("pdf", "PDF"),
              ("html", "srv_convert.pagina_web"), ("xps", "XPS"), ("xml", "srv_convert.foglio_xml_2003")],
    "slides": [("pptx", "PowerPoint"), ("ppt", "PowerPoint 97-2003"), ("pptm", "srv_convert.powerpoint_macro"), ("ppsx", "srv_convert.presentazione_proiettare"),
               ("potx", "srv_convert.modello_powerpoint"), ("odp", "srv_convert.opendocument_presentazione"), ("pdf", "PDF"), ("xps", "XPS"),
               ("png", "srv_convert.immagini_png_diapositiva"), ("jpg", "srv_convert.immagini_jpg_diapositiva"), ("mp4", "srv_convert.video_mp4"), ("rtf", "srv_convert.struttura_rtf")],
    "pdf": [("docx", "srv_convert.word_modificabile"), ("txt", "srv_convert.testo"), ("html", "srv_convert.pagina_web"), ("svg", "srv_convert.svg_pagina_file"),
            ("png", "srv_convert.immagini_png"), ("jpg", "srv_convert.immagini_jpg"), ("tiff", "srv_convert.tiff_multipagina"), ("xlsx", "srv_convert.excel_tabelle_pdf"), ("pdf", "srv_convert.pdf_ottimizzato"),
            ("dxf", "srv_convert.disegno_dxf"), ("dwg", "srv_convert.disegno_dwg")],
    "xps": [("pdf", "PDF"), ("png", "srv_convert.immagini_png"), ("txt", "srv_convert.testo")],
    "epub": [("pdf", "PDF"), ("txt", "srv_convert.testo"), ("html", "srv_convert.pagina_web"), ("docx", "Word")],
    "svg": [("png", "PNG"), ("jpg", "JPG"), ("webp", "WEBP"), ("tiff", "TIFF"), ("pdf", "PDF")],
    "html": [("pdf", "PDF"), ("docx", "Word"), ("odt", "srv_convert.opendocument_testo"), ("rtf", "RTF"), ("txt", "srv_convert.testo"), ("mht", "srv_convert.pagina_web_file")],
    "markdown": [("html", "srv_convert.pagina_web"), ("pdf", "PDF"), ("docx", "Word"), ("odt", "srv_convert.opendocument_testo"), ("txt", "srv_convert.testo")],
    "image": [(e, e.upper()) for e in ("png", "jpg", "webp", "tiff", "bmp", "gif", "ico", "avif", "heic", "jp2", "tga", "ppm", "pcx", "dds", "icns", "qoi")] + [("pdf", "PDF"),
              ("dxf", "srv_convert.disegno_dxf_scansione"), ("dwg", "srv_convert.disegno_dwg_scansione")],
    "media": [(e, e.upper()) for e in ("mp4", "mkv", "webm", "mov", "avi", "wmv", "flv", "mpg", "ts", "3gp", "gif")] + [(e, "srv_convert.solo_audio|" + e.upper()) for e in ("mp3", "wav", "flac", "m4a", "ogg", "opus")],
    "audio": [(e, e.upper()) for e in ("mp3", "wav", "flac", "m4a", "aac", "ogg", "opus", "wma", "aiff", "ac3")],
    "archive": [("zip", "ZIP"), ("7z", "7-Zip"), ("tar", "TAR"), ("tgz", "TAR.GZ"), ("txz", "TAR.XZ"), ("tbz2", "TAR.BZ2"), ("wim", "WIM")],
    "font": [("ttf", "TrueType"), ("otf", "OpenType"), ("woff", "WOFF"), ("woff2", "WOFF2")],
    "database": [("xlsx", "srv_convert.excel_tabella_foglio"), ("csv", "srv_convert.csv_tabelle_file_zip"), ("sql", "srv_convert.script_sql"), ("json", "JSON")],
    "text": [("txt", "srv_convert.testo"), ("pdf", "PDF"), ("docx", "Word"), ("html", "srv_convert.pagina_web")],
    "email": [("eml", "srv_convert.email"), ("html", "srv_convert.pagina_web"), ("txt", "srv_convert.testo"), ("pdf", "PDF")],
}

WORD_FMT = {"docx": 16, "doc": 0, "docm": 13, "dotx": 14, "rtf": 6, "odt": 23, "pdf": 17, "xps": 18, "txt": 7, "html": 10, "htm": 10, "mht": 9, "xml": 19}
EXCEL_FMT = {"xlsx": 51, "xlsm": 52, "xlsb": 50, "xls": 56, "csv": 62, "ods": 60, "html": 44, "txt": 42, "xltx": 54, "xml": 46}
PPT_FMT = {"pptx": 24, "ppt": 1, "pptm": 25, "ppsx": 28, "potx": 26, "odp": 35, "pdf": 32, "xps": 33, "png": 18, "jpg": 17, "rtf": 6, "mp4": 39}


def _label(l: str) -> str:
    if not l.startswith("srv_convert."):
        return l
    key, _, extra = l.partition("|")
    return tr(key) + extra


def targets(path: str) -> dict:
    ext = Path(path).suffix.lower().lstrip(".")
    fam = FAMILY_OF.get(ext)
    lst = [{"ext": e, "label": _label(l)} for e, l in TARGETS.get(fam, []) if e != ext or fam in ("pdf",)]
    if not _oda():  # DWG solo con ODA File Converter installato
        lst = [x for x in lst if x["ext"] != "dwg"]
    return {"ok": True, "family": fam, "targets": lst}


def _oda() -> bool:
    try:
        from .cadedit import find_oda
        return bool(find_oda())
    except Exception:  # noqa: BLE001
        return False


def convert(src: str, out: str, update=None, opts: dict | None = None) -> dict:
    update = update or (lambda **k: None)
    ext = Path(src).suffix.lower().lstrip(".")
    to = Path(out).suffix.lower().lstrip(".")
    fam = FAMILY_OF.get(ext)
    if not fam:
        raise ValueError(tr("srv_convert.nessuna_conversione_disponibile_i_file", ext))
    if to in ("dxf", "dwg") and fam in ("pdf", "image"):
        # PDF vettoriali, scansioni e immagini -> disegno (mbvisual/todrawing.py); opzioni predefinite: scala 1:1 della carta
        from . import todrawing
        return todrawing.convert(src, out, opts or {}, update=update)
    update(message=tr("srv_convert.conversione_corso"), progress=0.1)
    fn = {
        "word": _office_doc, "sheet": _office_sheet, "slides": _office_slides, "pdf": _pdf, "xps": _mupdf_doc, "epub": _mupdf_doc,
        "svg": _svg, "html": _html, "markdown": _markdown, "image": _image, "media": _media, "audio": _media, "archive": _archive,
        "font": _font, "database": _database, "text": _text, "email": _email,
    }[fam]
    res = fn(src, out, ext, to, update) or {}
    return {"ok": True, "path": out, **res}


# ---------------------------------------------------------------- Office (COM, poi LibreOffice)
def find_office(exe):
    """Office utilizzabile per l'automazione COM: solo su Windows (sul Mac si converte con LibreOffice)."""
    return _find_office(exe) if IS_WIN else None


def _com(progid):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx(progid)
    try:
        app.DisplayAlerts = False if progid != "PowerPoint.Application" else 1  # ppAlertsNone
    except Exception:  # noqa: BLE001
        pass
    return app


def _lo(src, out, to):
    exe = find_libreoffice()
    if not exe:
        raise RuntimeError(tr("srv_convert.questa_conversione_serve_microsoft_office"))
    import subprocess
    tmp = Path(tempfile.mkdtemp(dir=TEMP_DIR))
    profile = (TEMP_DIR / "lo-profile").as_uri()
    subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--convert-to", to, "--outdir", str(tmp), src],
                   capture_output=True, timeout=900, creationflags=NO_WINDOW)
    found = list(tmp.glob("*"))
    if not found:
        raise RuntimeError(tr("srv_convert.conversione_libreoffice_non_riuscita"))
    shutil.move(str(found[0]), out)


def _office_doc(src, out, ext, to, update):
    if to == "epub" or not find_office("winword.exe"):
        return _lo(src, out, to)
    word = _com("Word.Application")
    try:
        doc = word.Documents.Open(os.path.abspath(src), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False, Visible=False)
        update(message=tr("srv_convert.salvataggio_word"), progress=0.5)
        if to in WORD_FMT:
            doc.SaveAs2(os.path.abspath(out), FileFormat=WORD_FMT[to])
        else:
            raise ValueError(tr("srv_convert.word_non_salva_0", to))
        doc.Close(False)
    finally:
        word.Quit()


def _office_sheet(src, out, ext, to, update):
    if not find_office("excel.exe"):
        return _lo(src, out, to)
    xl = _com("Excel.Application")
    try:
        wb = xl.Workbooks.Open(os.path.abspath(src), ReadOnly=True, AddToMru=False)
        update(message=tr("srv_convert.salvataggio_excel"), progress=0.5)
        if to in ("pdf", "xps"):
            wb.ExportAsFixedFormat(0 if to == "pdf" else 1, os.path.abspath(out))
        elif to in EXCEL_FMT:
            wb.SaveAs(os.path.abspath(out), FileFormat=EXCEL_FMT[to])
        else:
            raise ValueError(tr("srv_convert.excel_non_salva_0", to))
        wb.Close(False)
    finally:
        xl.Quit()


def _office_slides(src, out, ext, to, update):
    if not find_office("powerpnt.exe"):
        return _lo(src, out, to)
    pp = _com("PowerPoint.Application")
    try:
        pres = pp.Presentations.Open(os.path.abspath(src), ReadOnly=True, Untitled=False, WithWindow=False)
        update(message=tr("srv_convert.salvataggio_powerpoint"), progress=0.5)
        if to == "mp4":
            pres.CreateVideo(os.path.abspath(out), True, 5, 1080, 30, 85)
            import time
            while pres.CreateVideoStatus == 1:  # in corso
                time.sleep(1)
        elif to in ("png", "jpg"):
            folder = Path(out).with_suffix("")
            folder.mkdir(parents=True, exist_ok=True)
            for i, slide in enumerate(pres.Slides, 1):
                slide.Export(str(folder / f"Diapositiva_{i:03d}.{to}"), to.upper(), 1920, 1080)
            pres.Close()
            return {"cartella": str(folder)}
        elif to in PPT_FMT:
            pres.SaveAs(os.path.abspath(out), PPT_FMT[to])
        else:
            raise ValueError(tr("srv_convert.powerpoint_non_salva_0", to))
        pres.Close()
    finally:
        pp.Quit()


# ---------------------------------------------------------------- PDF e documenti a pagine (PyMuPDF)
def _pages_to_images(doc, out, to, dpi=150):
    import pymupdf  # noqa: F401
    if len(doc) == 1:
        doc[0].get_pixmap(dpi=dpi).save(out)
        return {}
    stem, folder = Path(out).stem, Path(out).parent
    files = []
    for i, page in enumerate(doc, 1):
        f = folder / f"{stem}_pag{i:03d}.{to}"
        page.get_pixmap(dpi=dpi).save(str(f))
        files.append(str(f))
    return {"files": files}


def _pdf(src, out, ext, to, update):
    import pymupdf
    if to == "docx":
        from . import pdfedit
        pdfedit.to_word(src, out)
        return
    if to == "pdf":
        from . import pdfedit
        pdfedit.compress(src, out)
        return
    with pymupdf.open(src) as doc:
        return _mupdf_out(doc, out, to, update)


def _mupdf_out(doc, out, to, update):
    import pymupdf
    if to == "txt":
        Path(out).write_text("\n\f".join(p.get_text() for p in doc), encoding="utf-8")
    elif to == "html":
        Path(out).write_text("<!doctype html><meta charset='utf-8'>" + "".join(p.get_text("html") for p in doc), encoding="utf-8")
    elif to == "svg":
        stem, folder = Path(out).stem, Path(out).parent
        for i, p in enumerate(doc, 1):
            (Path(out) if len(doc) == 1 else folder / f"{stem}_pag{i:03d}.svg").write_text(p.get_svg_image(), encoding="utf-8")
    elif to in ("png", "jpg", "jpeg", "webp"):
        if to == "webp":
            from PIL import Image
            pix = doc[0].get_pixmap(dpi=150)
            Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(out, "WEBP", quality=90)
            return {}
        return _pages_to_images(doc, out, to)
    elif to in ("tif", "tiff"):
        from PIL import Image
        ims = []
        for p in doc:
            pix = p.get_pixmap(dpi=200)
            ims.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        ims[0].save(out, "TIFF", save_all=True, append_images=ims[1:], compression="tiff_lzw")
    elif to == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        n = 0
        for pi, p in enumerate(doc, 1):
            for ti, t in enumerate(p.find_tables().tables, 1):
                ws = wb.create_sheet(f"Pag{pi}_Tab{ti}"[:31])
                for row in t.extract():
                    ws.append([("" if c is None else c) for c in row])
                n += 1
        if not n:
            ws = wb.create_sheet(tr("srv_convert.testo"))
            for p in doc:
                for line in p.get_text().splitlines():
                    ws.append([line])
        wb.save(out)
        return {"tabelle": n}
    elif to == "pdf":
        pdf = pymupdf.open("pdf", doc.convert_to_pdf())
        pdf.save(out, garbage=3, deflate=True)
    elif to == "docx":
        tmp = str(Path(tempfile.mkdtemp(dir=TEMP_DIR)) / "tmp.pdf")
        pymupdf.open("pdf", doc.convert_to_pdf()).save(tmp)
        from . import pdfedit
        pdfedit.to_word(tmp, out)
    else:
        raise ValueError(tr("srv_convert.formato_0_non_disponibile", to))
    return {}


def _mupdf_doc(src, out, ext, to, update):
    import pymupdf
    with pymupdf.open(src) as doc:
        return _mupdf_out(doc, out, to, update)


def _svg(src, out, ext, to, update):
    import pymupdf
    data = Path(src).read_bytes()
    if ext == "svgz":
        import gzip
        data = gzip.decompress(data)
    doc = pymupdf.open("svg", data)
    if to == "pdf":
        pymupdf.open("pdf", doc.convert_to_pdf()).save(out)
        return
    pix = doc[0].get_pixmap(dpi=200, alpha=to in ("png", "webp", "tiff"))
    png = pix.tobytes("png")
    misc.save_image(png, out, 92)


# ---------------------------------------------------------------- HTML, Markdown, testo, email
def _html_to(html_text: str, out: str, to: str, base_dir: str | None = None):
    if to in ("html", "htm"):
        Path(out).write_text(html_text, encoding="utf-8")
        return
    if to == "txt":
        import re
        txt = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html_text, flags=re.S | re.I)
        txt = re.sub(r"<br\s*/?>|</p>|</div>|</h\d>|</li>|</tr>", "\n", txt, flags=re.I)
        Path(out).write_text(htmlmod.unescape(re.sub(r"<[^>]+>", "", txt)), encoding="utf-8")
        return
    tmp = Path(tempfile.mkdtemp(dir=TEMP_DIR)) / "doc.html"
    tmp.write_text(html_text, encoding="utf-8")
    if find_office("winword.exe") and to in WORD_FMT and to != "pdf":
        return _office_doc(str(tmp), out, "html", to, lambda **k: None)
    if to == "pdf":  # PyMuPDF: impaginazione A4 immediata (Word impiegherebbe decine di secondi)
        import pymupdf
        story = pymupdf.Story(html=html_text, archive=base_dir)
        w = pymupdf.DocumentWriter(out)
        more = True
        while more:
            dev = w.begin_page(pymupdf.paper_rect("a4"))
            more, _ = story.place(pymupdf.paper_rect("a4") + (50, 50, -50, -50))
            story.draw(dev)
            w.end_page()
        w.close()
        return
    if to == "docx":
        from . import office
        office.save_docx_from_html(html_text, out)
        return
    _lo(str(tmp), out, to)


def _html(src, out, ext, to, update):
    if to == "mht" and find_office("winword.exe"):
        return _office_doc(src, out, ext, to, update)
    raw = Path(src).read_bytes()
    text = raw.decode("utf-8", "replace")
    _html_to(text, out, to, str(Path(src).parent))


def _markdown(src, out, ext, to, update):
    import markdown
    body = markdown.markdown(Path(src).read_text(encoding="utf-8", errors="replace"), extensions=["tables", "fenced_code", "toc", "sane_lists"])
    page = ("<!doctype html><html><head><meta charset='utf-8'><style>body{font-family:Calibri,Arial,sans-serif;max-width:820px;margin:auto;line-height:1.45}"
            "table{border-collapse:collapse}td,th{border:1px solid #999;padding:4px 8px}code,pre{font-family:Consolas,monospace;background:#f4f4f4}</style>"
            f"</head><body>{body}</body></html>")
    if to == "txt":
        shutil.copyfile(src, out)
        return
    _html_to(page, out, to, str(Path(src).parent))


def _text(src, out, ext, to, update):
    raw = misc.read_text(src)
    text = raw.get("text", "") if isinstance(raw, dict) else ""
    if to == "txt":
        Path(out).write_text(text, encoding="utf-8")
        return
    page = f"<!doctype html><meta charset='utf-8'><pre style='font-family:Consolas,monospace;font-size:10pt;white-space:pre-wrap'>{htmlmod.escape(text)}</pre>"
    _html_to(page, out, to)


def _email(src, out, ext, to, update):
    if to == "eml" and ext == "eml":
        shutil.copyfile(src, out)
        return
    if to == "eml":  # Outlook .msg -> .eml standard (testo, HTML e allegati)
        import mimetypes
        from email.message import EmailMessage
        from email.policy import SMTP
        m = (misc.read_email(src).get("messaggi") or [{}])[0]
        msg = EmailMessage()
        for k, h in (("da", "From"), ("a", "To"), ("cc", "Cc"), ("data", "Date"), ("oggetto", "Subject")):
            if m.get(k):
                msg[h] = " ".join(str(m[k]).split())  # intestazioni su una riga
        msg.set_content(m.get("testo") or "")
        if m.get("html"):
            msg.add_alternative(m["html"], subtype="html")
        for a in m.get("allegati") or []:
            try:
                data = Path(a["path"]).read_bytes()
            except OSError:
                continue
            mt = (a.get("tipo") or mimetypes.guess_type(a["nome"])[0] or "application/octet-stream").split("/", 1)
            msg.add_attachment(data, maintype=mt[0], subtype=mt[1] if len(mt) > 1 else "octet-stream", filename=a["nome"])
        Path(out).write_bytes(msg.as_bytes(policy=SMTP))
        return
    parts = []
    for m in misc.read_email(src).get("messaggi", []):
        head = "".join(f"<b>{htmlmod.escape(k.capitalize())}:</b> {htmlmod.escape(str(m.get(k) or ''))}<br>" for k in ("da", "a", "cc", "data", "oggetto") if m.get(k))
        atts = ", ".join(a["nome"] for a in m.get("allegati") or [])
        body = m.get("html") or f"<pre style='white-space:pre-wrap'>{htmlmod.escape(m.get('testo') or '')}</pre>"
        parts.append(f"<div style='font-family:Calibri,Arial'>{head}{'<b>Allegati:</b> ' + htmlmod.escape(atts) + '<br>' if atts else ''}<hr>{body}</div>")
    _html_to("<!doctype html><meta charset='utf-8'>" + "<hr style='border:2px solid #888'>".join(parts), out, to)


# ---------------------------------------------------------------- immagini, audio/video
def _image(src, out, ext, to, update):
    Image = misc._pil()
    if ext in misc.RAW_EXTS:
        png = Path(misc.image_to_png(src)["path"]).read_bytes()
    else:
        with Image.open(src) as im:
            buf = io.BytesIO()
            frame = im.convert("RGBA") if im.mode not in ("RGB", "RGBA", "L", "LA") else im.copy()
            frame.save(buf, "PNG")
            png = buf.getvalue()
    if to == "pdf":
        im = Image.open(io.BytesIO(png)).convert("RGB")
        im.save(out, "PDF", resolution=150)
        return
    misc.save_image(png, out, 92)


def _media(src, out, ext, to, update):
    from . import media
    media.edit(src, out, update=update)


# ---------------------------------------------------------------- archivi, font, database
def _archive(src, out, ext, to, update):
    from . import archive
    from .paths import find_7zip
    tmp = Path(tempfile.mkdtemp(dir=TEMP_DIR))
    update(message=tr("srv_convert.estrazione"), progress=0.2)
    archive.extract(src, None, str(tmp))
    update(message=tr("srv_convert.compressione"), progress=0.6)
    exe = find_7zip()
    import subprocess
    if os.path.exists(out):
        os.remove(out)
    if to in ("tgz", "txz", "tbz2"):
        tar = str(tmp.parent / (tmp.name + ".tar"))
        subprocess.run([exe, "a", "-ttar", tar, str(tmp / "*")], capture_output=True, creationflags=NO_WINDOW, check=True)
        kind = {"tgz": "gzip", "txz": "xz", "tbz2": "bzip2"}[to]
        subprocess.run([exe, "a", f"-t{kind}", out, tar], capture_output=True, creationflags=NO_WINDOW, check=True)
    else:
        kind = {"zip": "zip", "7z": "7z", "tar": "tar", "wim": "wim"}[to]
        subprocess.run([exe, "a", f"-t{kind}", "-mx=7", out, str(tmp / "*")], capture_output=True, creationflags=NO_WINDOW, check=True)


def _font(src, out, ext, to, update):
    from fontTools.ttLib import TTFont
    f = TTFont(src)
    f.flavor = {"woff": "woff", "woff2": "woff2"}.get(to)
    f.save(out)


def _database(src, out, ext, to, update):
    con = sqlite3.connect(f"file:{src}?mode=ro", uri=True)
    try:
        tables = [r[0] for r in con.execute("select name from sqlite_master where type='table' and name not like 'sqlite_%' order by name")]
        if to == "sql":
            with open(out, "w", encoding="utf-8") as f:
                for line in con.iterdump():
                    f.write(line + "\n")
        elif to == "json":
            data = {}
            for t in tables:
                cur = con.execute(f'select * from "{t}"')
                cols = [c[0] for c in cur.description]
                data[t] = [dict(zip(cols, [misc._cell(v) for v in r])) for r in cur]
            Path(out).write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        elif to == "xlsx":
            import openpyxl
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            for t in tables:
                ws = wb.create_sheet(t[:31])
                cur = con.execute(f'select * from "{t}"')
                ws.append([c[0] for c in cur.description])
                for r in cur:
                    ws.append([misc._cell(v) for v in r])
            wb.save(out)
        elif to == "csv":
            with zipfile.ZipFile(out if out.endswith(".zip") else out, "w", zipfile.ZIP_DEFLATED) as z:
                for t in tables:
                    buf = io.StringIO()
                    cur = con.execute(f'select * from "{t}"')
                    wr = csv.writer(buf, delimiter=";")
                    wr.writerow([c[0] for c in cur.description])
                    for r in cur:
                        wr.writerow([misc._cell(v) for v in r])
                    z.writestr(f"{t}.csv", "﻿" + buf.getvalue())
        else:
            raise ValueError(tr("srv_convert.formato_0_non_disponibile_i", to))
    finally:
        con.close()
