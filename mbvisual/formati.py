"""Formati aggiuntivi (un'unica API /api/fx?op=...): DICOM, certificati, dati scientifici/tabellari, Apple iWork,
Revit, piu' le esportazioni in PDF/Excel usate da calendari, contatti e file di plotter (letti nel browser).

Ogni operazione ritorna un dizionario (risposta JSON) oppure (byte, intestazioni) per i dati binari.
"""
from __future__ import annotations

import json
import os
from pathlib import Path


def handle(op: str, path: str | None, qs: dict, body: bytes):
    q = lambda k, d=None: (qs.get(k) or [d])[0]  # noqa: E731
    js = lambda: json.loads(body.decode("utf-8")) if body else {}  # noqa: E731
    # ---------------------------------------------------------------- DICOM
    if op == "dicom_open":
        from . import fx_dicom
        return fx_dicom.open_series(path, q("uid"))
    if op == "dicom_frame":
        from . import fx_dicom
        data, hdr = fx_dicom.frame(path, int(q("frame", "0")))
        return data, hdr
    if op == "dicom_tags":
        from . import fx_dicom
        return fx_dicom.tags(path)
    # ---------------------------------------------------------------- certificati
    if op == "cert_read":
        from . import fx_cert
        return fx_cert.read(path, js().get("password"))
    if op == "cert_export":
        from . import fx_cert
        b = js()
        return fx_cert.export(b["pem"], b["out"], b.get("indice"))
    # ---------------------------------------------------------------- dati scientifici e tabellari
    if op == "sci_open":
        from . import fx_sci
        return fx_sci.open_file(path)
    if op == "sci_data":
        from . import fx_sci
        b = js()
        return fx_sci.data(path, b.get("nome") or "", int(b.get("da", 0)), int(b.get("n", 200)), b.get("fetta"))
    if op == "sci_export":
        from . import fx_sci
        b = js()
        return fx_sci.export(path, b.get("nome") or "", b["out"], b.get("fetta"))
    # ---------------------------------------------------------------- Apple iWork
    if op == "iwork_read":
        from . import fx_iwork
        return fx_iwork.read(path)
    if op == "iwork_part":
        from . import fx_iwork
        return fx_iwork.part(path, q("nome"))
    if op == "iwork_page":
        from . import fx_iwork
        return fx_iwork.page_png(path, int(q("n", "0")))
    if op == "iwork_export":
        from . import fx_iwork
        b = js()
        return fx_iwork.export(path, b["out"])
    # ---------------------------------------------------------------- Revit
    if op == "revit_read":
        from . import fx_revit
        return fx_revit.read(path)
    if op == "revit_preview":
        from . import fx_revit
        return fx_revit.preview(path)
    if op == "revit_export":
        from . import fx_revit
        return fx_revit.export(path, js()["out"])
    # ---------------------------------------------------------------- esportazioni comuni
    if op == "html_pdf":
        b = js()
        return html_pdf(b["html"], b["out"], b.get("formato", "A4"), bool(b.get("orizzontale")))
    if op == "svg_pdf":
        b = js()
        return svg_pdf(b["svg"], b["out"])
    if op == "xlsx":
        b = js()
        return rows_xlsx(b["fogli"], b["out"])
    raise ValueError(f"Operazione sconosciuta: {op}")


def html_pdf(html: str, out: str, formato="A4", orizzontale=False) -> dict:
    """HTML semplice (tabelle, titoli) -> PDF con PyMuPDF Story, a piu' pagine."""
    import pymupdf as fitz
    rect = fitz.paper_rect(formato.lower() + ("-l" if orizzontale else ""))
    story = fitz.Story(html=html)
    writer = fitz.DocumentWriter(out)
    where = rect + (36, 36, -36, -36)
    more = True
    while more:
        dev = writer.begin_page(rect)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return {"ok": True, "path": out}


def svg_pdf(svg: str, out: str) -> dict:
    """Disegno SVG (vettoriale) -> PDF di una pagina grande quanto il disegno."""
    import pymupdf as fitz
    src = fitz.open("svg", svg.encode("utf-8"))
    pdf = fitz.open("pdf", src.convert_to_pdf())
    pdf.save(out, garbage=3, deflate=True)
    return {"ok": True, "path": out}


def rows_xlsx(fogli: list, out: str) -> dict:
    """[{nome, colonne, righe}] -> Excel con intestazioni in grassetto e colonne adattate."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for f in fogli:
        ws = wb.create_sheet(_sheet_name(f.get("nome") or "Foglio", wb))
        cols = f.get("colonne") or []
        if cols:
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
        for r in f.get("righe") or []:
            ws.append([_cell(v) for v in r])
        for i, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(r[i - 1])) for r in (f.get("righe") or [])[:500] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, max(8, width + 2))
    if not wb.worksheets:
        wb.create_sheet("Foglio1")
    wb.save(out)
    return {"ok": True, "path": out}


def _cell(v):
    if isinstance(v, (int, float, bool)) or v is None:
        return v
    s = str(v)
    return s[:32000]


def _sheet_name(name: str, wb) -> str:
    bad = '[]:*?/\\'
    base = "".join("_" if ch in bad else ch for ch in name)[:31] or "Foglio"
    n, k = base, 2
    while n in wb.sheetnames:
        suffix = f" ({k})"
        n = base[:31 - len(suffix)] + suffix
        k += 1
    return n


def write_csv(path: str, cols: list, rows) -> None:
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        if cols:
            w.writerow(cols)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])


def ext_of(p: str) -> str:
    return Path(p).suffix.lower().lstrip(".")


def ensure_dir(p: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(p)), exist_ok=True)
