"""Documenti Office: PowerPoint (lettura/modifica testi), Word da HTML, conversioni con LibreOffice."""
import base64
import hashlib
import io
import os
import subprocess
from pathlib import Path

from .paths import TEMP_DIR, find_libreoffice

NO_WINDOW = 0x08000000
EMU_PER_PX = 9525


# ---------------------------------------------------------------- LibreOffice
def lo_convert(path: str, to: str = "pdf", update=None) -> dict:
    exe = find_libreoffice()
    if not exe:
        raise RuntimeError("LibreOffice non e' installato su questo PC. E' gratuito: https://www.libreoffice.org/download/download/")
    st = os.stat(path)
    key = hashlib.sha1(f"{path}|{st.st_size}|{st.st_mtime}|{to}".encode()).hexdigest()[:16]
    out_dir = TEMP_DIR / "office" / key
    target = out_dir / (Path(path).stem + "." + to.split(":")[0])
    if target.exists():
        return {"path": str(target)}
    out_dir.mkdir(parents=True, exist_ok=True)
    if update:
        update(message="Conversione con LibreOffice", progress=0.1)
    profile = (TEMP_DIR / "lo-profile").as_uri()
    r = subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--convert-to", to,
                        "--outdir", str(out_dir), path], capture_output=True, timeout=600, creationflags=NO_WINDOW)
    if not target.exists():
        found = list(out_dir.glob("*"))
        if found:
            return {"path": str(found[0])}
        raise RuntimeError("Conversione non riuscita: " + r.stderr.decode("utf-8", "replace")[-400:])
    return {"path": str(target)}


# ---------------------------------------------------------------- PowerPoint
def _color(fill):
    try:
        if fill.type == 1:  # MSO_FILL.SOLID
            return "#" + str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def _paragraphs(tf):
    out = []
    for p in tf.paragraphs:
        size = None
        bold = italic = False
        color = None
        for r in p.runs:
            f = r.font
            if f.size and not size:
                size = f.size.pt
            bold = bold or bool(f.bold)
            italic = italic or bool(f.italic)
            try:
                if f.color and f.color.type is not None and f.color.rgb is not None and not color:
                    color = "#" + str(f.color.rgb)
            except Exception:
                pass
        align = None
        if p.alignment is not None:
            align = {1: "left", 2: "center", 3: "right", 4: "justify"}.get(int(p.alignment), None)
        out.append({"text": "".join(r.text for r in p.runs) if p.runs else p.text, "level": p.level,
                    "size": size, "bold": bold, "italic": italic, "color": color, "align": align})
    return out


def _shape_json(sh, ox=0, oy=0):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    items = []
    try:
        x, y, w, h = (sh.left or 0) + ox, (sh.top or 0) + oy, sh.width or 0, sh.height or 0
    except Exception:
        x = y = w = h = 0
    base = {"id": sh.shape_id, "name": sh.name, "x": x / EMU_PER_PX, "y": y / EMU_PER_PX,
            "w": w / EMU_PER_PX, "h": h / EMU_PER_PX, "rot": getattr(sh, "rotation", 0) or 0}
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.GROUP:
        for child in sh.shapes:
            items.extend(_shape_json(child, 0, 0))
        return items
    try:
        img = sh.image
        base.update(kind="image", src=f"data:{img.content_type};base64," + base64.b64encode(img.blob).decode())
        items.append(base)
        return items
    except Exception:
        pass
    if getattr(sh, "has_table", False) and sh.has_table:
        rows = []
        for r in sh.table.rows:
            rows.append([_paragraphs(c.text_frame) for c in r.cells])
        base.update(kind="table", rows=rows)
        items.append(base)
        return items
    fill = None
    try:
        fill = _color(sh.fill)
    except Exception:
        pass
    line = None
    try:
        if sh.line and sh.line.fill.type == 1:
            line = "#" + str(sh.line.color.rgb)
    except Exception:
        pass
    base.update(kind="shape", fill=fill, line=line)
    try:
        if sh.is_placeholder:
            base["ph"] = {1: "title", 3: "ctrTitle", 4: "subTitle", 2: "body", 7: "body"}.get(int(sh.placeholder_format.type), "other")
    except Exception:
        pass
    if sh.has_text_frame:
        base["paragraphs"] = _paragraphs(sh.text_frame)
        try:
            base["anchor"] = {1: "top", 3: "middle", 4: "bottom"}.get(int(sh.text_frame.vertical_anchor), None) if sh.text_frame.vertical_anchor else None
        except Exception:
            base["anchor"] = None
    items.append(base)
    return items


def read_pptx(path: str) -> dict:
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, s in enumerate(prs.slides):
        shapes = []
        for sh in s.shapes:
            shapes.extend(_shape_json(sh))
        bg = None
        try:
            bg = _color(s.background.fill)
        except Exception:
            pass
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        slides.append({"index": i, "shapes": shapes, "background": bg, "notes": notes})
    return {"ok": True, "width": prs.slide_width / EMU_PER_PX, "height": prs.slide_height / EMU_PER_PX, "slides": slides}


def _set_text(tf, texts: list[str]):
    paras = list(tf.paragraphs)
    for i, t in enumerate(texts):
        if i < len(paras):
            p = paras[i]
            runs = list(p.runs)
            if runs:
                runs[0].text = t
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
            else:
                p.text = t
        else:
            p = tf.add_paragraph()
            if paras and paras[-1].runs:
                p.level = paras[-1].level
            r = p.add_run()
            r.text = t
            if paras and paras[-1].runs:
                src = paras[-1].runs[0].font
                r.font.size, r.font.bold, r.font.italic = src.size, src.bold, src.italic
    for p in paras[len(texts):]:
        p._p.getparent().remove(p._p)


def _find_shape(shapes, sid):
    for sh in shapes:
        if sh.shape_id == sid:
            return sh
        if sh.shape_type == 6:  # group
            f = _find_shape(sh.shapes, sid)
            if f is not None:
                return f
    return None


def save_pptx(path: str, out_path: str, edits: list[dict]) -> dict:
    from pptx import Presentation
    prs = Presentation(path)
    slides = list(prs.slides)
    for e in edits:
        sh = _find_shape(slides[e["slide"]].shapes, e["shape"])
        if sh is None:
            continue
        if "cell" in e:
            r, c = e["cell"]
            _set_text(sh.table.cell(r, c).text_frame, e["paragraphs"])
        elif sh.has_text_frame:
            _set_text(sh.text_frame, e["paragraphs"])
    for n in sorted({e["delete_slide"] for e in edits if "delete_slide" in e}, reverse=True):
        sld_id = prs.slides._sldIdLst[n]
        prs.part.drop_rel(sld_id.rId)
        prs.slides._sldIdLst.remove(sld_id)
    prs.save(out_path)
    return {"ok": True, "path": out_path}


# ---------------------------------------------------------------- Word
def save_docx_from_html(html: str, out_path: str, title: str = "") -> dict:
    from html2docx import html2docx
    buf = html2docx(html, title=title or Path(out_path).stem)
    data = buf.getvalue() if isinstance(buf, io.BytesIO) else bytes(buf)
    Path(out_path).write_bytes(data)
    return {"ok": True, "path": out_path}
