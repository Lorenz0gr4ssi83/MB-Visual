"""Cartella di prova per la ricerca nel progetto (mbvisual/search.py): file di ogni tipo con parole note + tante copie.

    python tests/make_search_samples.py <cartella> [--copie 3000] [--dwg file.dwg ...] [--ifc file.ifc ...]

Ogni file contiene un codice univoco (es. "ZEBRA-PDF-01") per verificare che la ricerca trovi il punto giusto.
Le copie (sottocartelle Lotto_NN) servono a misurare i tempi su migliaia di file.
"""
import argparse
import ctypes
import email.message
import io
import os
import random
import shutil
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WORDS = ("calcestruzzo armatura trave pilastro solaio fondazione copertura serramento impianto elettrico idraulico antincendio "
         "relazione tecnica computo metrico capitolato collaudo cantiere sicurezza progetto esecutivo verifica sismica carpenteria "
         "tubazione valvola pompa serbatoio silo nastro trasportatore quadro cavo canalina illuminazione ventilazione").split()


def pdfs(d: Path):
    import pymupdf
    doc = pymupdf.open()
    for i in range(1, 6):
        p = doc.new_page()
        text = f"Relazione strutturale - pagina {i}\n" + " ".join(random.sample(WORDS, 12))
        if i == 3:
            text += "\nLa trave ZEBRA-PDF-03 in acciaio S355 ha luce di 7,50 m. Resistenza al fuoco REI 120."
        p.insert_text((50, 80), text, fontsize=11)
    doc.set_metadata({"title": "Relazione di calcolo strutture", "author": "Ing. Mario Rossi", "keywords": "strutture, acciaio"})
    doc.save(d / "Relazione_calcolo.pdf")
    # scansione senza testo (solo immagine)
    doc = pymupdf.open()
    p = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), 0)
    pix.clear_with(200)
    p.insert_image(p.rect, pixmap=pix)
    doc.save(d / "Scansione_verbale.pdf")
    # protetto da password
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 80), "Contenuto riservato ZEBRA-SEGRETO")
    doc.save(d / "Offerta_riservata.pdf", encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    # PDF con testo OCR invisibile sotto un'immagine (come "Rendi ricercabile")
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_image(p.rect, pixmap=pix)
    p.insert_text((50, 80), "Verbale di sopralluogo ZEBRA-OCR-01 del cantiere di Brescia", render_mode=3)
    doc.save(d / "Verbale_OCR.pdf")


def office(d: Path):
    import docx
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    w = docx.Document()
    w.core_properties.title = "Capitolato speciale d'appalto"
    w.core_properties.author = "Studio MB"
    w.sections[0].header.paragraphs[0].text = "Intestazione commessa ZEBRA-INTEST"
    w.add_heading("Capo 1 - Opere edili", 1)
    w.add_paragraph("Le murature saranno in blocchi di laterizio. " + " ".join(random.sample(WORDS, 10)))
    w.add_heading("Capo 2 - Serramenti", 1)
    w.add_paragraph("I serramenti esterni in alluminio a taglio termico codice ZEBRA-DOCX-02 con vetrocamera basso emissivo.")
    w.save(d / "Capitolato.docx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Computo"
    ws.append(["Codice", "Descrizione", "Quantita'", "Prezzo"])
    for i in range(2, 60):
        ws.append([f"A.{i:03d}", " ".join(random.sample(WORDS, 4)), i * 1.5, i * 10])
    ws["B37"] = "Fornitura di pompa centrifuga ZEBRA-XLSX-B37 portata 30 m3/h"
    ws2 = wb.create_sheet("Riepilogo")
    ws2["C4"] = "Totale lavori ZEBRA-XLSX-RIEP"
    wb.save(d / "Computo_metrico.xlsx")
    pr = Presentation()
    for i in range(1, 5):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = f"Avanzamento lavori - fase {i}"
        s.placeholders[1].text = " ".join(random.sample(WORDS, 8)) + (" consegna ZEBRA-PPTX-04" if i == 4 else "")
        if i == 2:
            s.notes_slide.notes_text_frame.text = "Nota del relatore ZEBRA-NOTE-02"
    pr.save(d / "Presentazione_SAL.pptx")
    # OpenDocument minimali scritti a mano
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
          'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:dc="http://purl.org/dc/elements/1.1/" '
          'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0"')
    meta = f'<?xml version="1.0"?><office:document-meta {ns}><office:meta><dc:title>Verbale riunione</dc:title></office:meta></office:document-meta>'
    for name, mt, body in (
        ("Verbale.odt", "application/vnd.oasis.opendocument.text",
         '<office:text><text:h>Verbale</text:h><text:p>Presenti: direttore lavori. Decisione ZEBRA-ODT-01 sulla variante.</text:p></office:text>'),
        ("Tabella.ods", "application/vnd.oasis.opendocument.spreadsheet",
         '<office:spreadsheet><table:table table:name="Dati"><table:table-row><table:table-cell><text:p>x</text:p></table:table-cell></table:table-row>'
         '<table:table-row><table:table-cell table:number-columns-repeated="2"/><table:table-cell><text:p>ZEBRA-ODS-C2</text:p></table:table-cell></table:table-row>'
         '</table:table></office:spreadsheet>')):
        with zipfile.ZipFile(d / name, "w") as z:
            z.writestr("mimetype", mt, compress_type=zipfile.ZIP_STORED)
            z.writestr("content.xml", f'<?xml version="1.0"?><office:document-content {ns}><office:body>{body}</office:body></office:document-content>')
            z.writestr("meta.xml", meta)
    (d / "Lettera.rtf").write_text(r"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 Spett.le ditta, con la presente si comunica la sospensione ZEBRA-RTF-01 dei lavori.\par Cordiali saluti}", encoding="latin-1")


def office_legacy(d: Path):
    """.doc e .xls veri (via Word/Excel se installati)."""
    try:
        import win32com.client
    except ImportError:
        return
    try:
        wd = win32com.client.DispatchEx("Word.Application")
        wd.Visible = False
        doc = wd.Documents.Add()
        doc.Content.Text = "Documento Word 97 con il codice ZEBRA-DOC-97 per la prova di ricerca nel formato binario."
        doc.SaveAs2(str(d / "Vecchio_documento.doc"), FileFormat=0)
        doc.Close(False)
        wd.Quit()
    except Exception as e:  # noqa: BLE001
        print("Word non disponibile:", e)
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
        xl.Visible = False
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Add()
        wb.Worksheets(1).Range("D5").Value = "Voce ZEBRA-XLS-97 del vecchio foglio"
        wb.SaveAs(str(d / "Vecchio_foglio.xls"), FileFormat=56)
        wb.Close(False)
        xl.Quit()
    except Exception as e:  # noqa: BLE001
        print("Excel non disponibile:", e)


def texts(d: Path):
    (d / "note_cantiere.txt").write_text("\n".join(f"riga {i}: " + " ".join(random.sample(WORDS, 6)) for i in range(1, 300)) +
                                         "\nriga 300: attenzione alla perdita ZEBRA-TXT-300 sul serbatoio\n", encoding="utf-8")
    (d / "LEGGIMI.md").write_text("# Progetto\n\nIstruzioni per la **consegna** degli elaborati ZEBRA-MD-01.\n", encoding="utf-8")
    (d / "elenco.csv").write_text("codice;descrizione\nP-01;Pompa ZEBRA-CSV-01\n", encoding="cp1252")
    (d / "pagina.html").write_text("<html><head><title>Scheda tecnica</title></head><body><p>Valvola a sfera ZEBRA-HTML-01</p><script>var x='nascosto';</script></body></html>", encoding="utf-8")
    (d / "calcolo.py").write_text("# script di calcolo\nCARICO = 'ZEBRA-PY-01'\nprint(CARICO)\n", encoding="utf-8")
    (d / "Città_perché.txt").write_text("Parole accentate: perché città qualità ZEBRA-ACCENTI", encoding="utf-8")


def emails(d: Path, pdf: Path, docxf: Path):
    m = email.message.EmailMessage()
    m["From"] = "Impresa Bianchi <ufficio@impresabianchi.it>"
    m["To"] = "direzione.lavori@example.com"
    m["Subject"] = "Invio documentazione ZEBRA-EML-OGGETTO"
    m["Date"] = "Mon, 12 Oct 2026 10:00:00 +0200"
    m.set_content("Buongiorno, in allegato la relazione e il capitolato aggiornati.\nCordiali saluti")
    m.add_attachment(pdf.read_bytes(), maintype="application", subtype="pdf", filename="Relazione_allegata.pdf")
    m.add_attachment(docxf.read_bytes(), maintype="application", subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="Capitolato_allegato.docx")
    (d / "Invio_documentazione.eml").write_bytes(bytes(m))


def cad(d: Path):
    import ezdxf
    doc = ezdxf.new("R2018")
    doc.layers.add("TESTI")
    doc.layers.add("ZEBRA-LAYER-ARREDI")
    msp = doc.modelspace()
    msp.add_line((0, 0), (100, 0))
    msp.add_text("SOGGIORNO ZEBRA-DXF-TEXT", dxfattribs={"layer": "TESTI"}).set_placement((10, 10))
    msp.add_mtext("Nota:\\Pmuro portante {\\fArial|b1;ZEBRA-DXF-MTEXT} da non demolire", dxfattribs={"layer": "TESTI"}).set_location((20, 20))
    blk = doc.blocks.new("CARTIGLIO")
    blk.add_attdef("TITOLO", (0, 0), dxfattribs={"height": 2.5})
    ins = msp.add_blockref("CARTIGLIO", (200, 0))
    ins.add_attrib("TITOLO", "Pianta piano terra ZEBRA-DXF-ATTRIB", (200, 0))
    doc.saveas(d / "Pianta_PT.dxf")


def ifc(d: Path):
    (d / "Edificio.ifc").write_text("""ISO-10303-21;
HEADER;
FILE_DESCRIPTION(('ViewDefinition [CoordinationView]'),'2;1');
FILE_NAME('Edificio.ifc','2026-10-22T10:00:00',('MB'),('MB Engineering'),'','MB Visual prova','');
FILE_SCHEMA(('IFC2X3'));
ENDSEC;
DATA;
#1=IFCPROJECT('0YvctVUKr0kugbFTf53O9L',$,'Progetto prova',$,$,$,$,$,$);
#2=IFCBUILDING('1YvctVUKr0kugbFTf53O9L',$,'Edificio A',$,$,$,$,'Palazzina uffici',.ELEMENT.,$,$,$);
#3=IFCBUILDINGSTOREY('2YvctVUKr0kugbFTf53O9L',$,'Piano terra',$,$,$,$,$,.ELEMENT.,0.);
#10=IFCWALLSTANDARDCASE('3YvctVUKr0kugbFTf53O9L',$,'Muro esterno 30 cm',$,'Muro:Laterizio',$,$,'ZEBRA-IFC-TAG');
#11=IFCDOOR('4YvctVUKr0kugbFTf53O9L',$,'Porta REI 60',$,$,$,$,'P01',2.1,0.9);
#20=IFCPROPERTYSINGLEVALUE('FireRating',$,IFCLABEL('REI 120 ZEBRA-IFC-PSET'),$);
#21=IFCPROPERTYSINGLEVALUE('IsExternal',$,IFCBOOLEAN(.T.),$);
#22=IFCPROPERTYSET('5YvctVUKr0kugbFTf53O9L',$,'Pset_WallCommon',$,(#20,#21));
#23=IFCRELDEFINESBYPROPERTIES('6YvctVUKr0kugbFTf53O9L',$,$,$,(#10),#22);
#30=IFCMATERIAL('Laterizio porizzato ZEBRA-IFC-MAT');
#31=IFCRELASSOCIATESMATERIAL('7YvctVUKr0kugbFTf53O9L',$,$,$,(#10),#30);
#40=IFCRELCONTAINEDINSPATIALSTRUCTURE('8YvctVUKr0kugbFTf53O9L',$,$,$,(#10,#11),#3);
#41=IFCRELAGGREGATES('9YvctVUKr0kugbFTf53O9L',$,$,$,#2,(#3));
#50=IFCQUANTITYAREA('NetSideArea',$,$,12.5);
#51=IFCELEMENTQUANTITY('AYvctVUKr0kugbFTf53O9L',$,'Qto_DoorBaseQuantities',$,$,(#50));
#52=IFCRELDEFINESBYPROPERTIES('BYvctVUKr0kugbFTf53O9L',$,$,$,(#11),#51);
#60=IFCWALL('CYvctVUKr0kugbFTf53O9L',$,'Parete con \\X2\\00E8\\X0\\ accento',$,$,$,$,$);
ENDSEC;
END-ISO-10303-21;
""", encoding="latin-1")


def images(d: Path):
    from PIL import Image
    im = Image.new("RGB", (64, 48), (120, 160, 200))
    ex = Image.Exif()
    ex[0x010F] = "Canon"
    ex[0x0110] = "EOS R6"
    ex[0x013B] = "Fotografo ZEBRA-EXIF-ARTIST"
    ex[0x010E] = "Sopralluogo facciata nord"
    ex[0x9C9E] = "cantiere;facciata;ZEBRA-EXIF-KEYWORD".encode("utf-16-le") + b"\x00\x00"
    im.save(d / "Foto_sopralluogo.jpg", exif=ex)
    from PIL import PngImagePlugin
    info = PngImagePlugin.PngInfo()
    info.add_text("Description", "Schema impianto ZEBRA-PNG-TEXT")
    im.save(d / "Schema.png", pnginfo=info)


def archives(d: Path):
    with zipfile.ZipFile(d / "Consegna.zip", "w") as z:
        z.writestr("Tavole/ZEBRA-ZIP-TAVOLA-01.pdf", b"%PDF-1.4")
        z.writestr("Relazioni/relazione_geologica.docx", b"x")
    s = ROOT / "tests" / "samples" / "contratto_firmato.pdf.p7m"
    if s.exists():
        shutil.copy(s, d / "Contratto_firmato.pdf.p7m")


def cloud_only(d: Path):
    """File che sembra "solo su OneDrive" (attributo OFFLINE): l'indice deve leggere solo nome e proprieta'."""
    p = d / "Solo_online_ZEBRA-NUVOLA-NOME.pdf"
    p.write_text("contenuto che non deve finire nell'indice ZEBRA-NUVOLA-CONTENUTO")
    FILE_ATTRIBUTE_OFFLINE = 0x1000
    ctypes.windll.kernel32.SetFileAttributesW(str(p), FILE_ATTRIBUTE_OFFLINE)


def local_copy(src: str, dst: Path):
    st = os.stat(src)
    if getattr(st, "st_file_attributes", 0) & (0x400000 | 0x40000 | 0x1000):
        print("saltato (solo online):", src)
        return
    shutil.copy2(src, dst / Path(src).name)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dest")
    ap.add_argument("--copie", type=int, default=0)
    ap.add_argument("--dwg", nargs="*", default=[])
    ap.add_argument("--ifc", nargs="*", default=[])
    a = ap.parse_args()
    random.seed(7)
    d = Path(a.dest)
    if d.exists():
        shutil.rmtree(d)
    base = d / "Commessa_2026-101"
    for sub in ("01_Relazioni", "02_Computi", "03_Disegni", "04_Corrispondenza", "05_Foto", "06_BIM", "07_Varie"):
        (base / sub).mkdir(parents=True)
    pdfs(base / "01_Relazioni")
    office(base / "02_Computi")
    office_legacy(base / "02_Computi")
    texts(base / "07_Varie")
    emails(base / "04_Corrispondenza", base / "01_Relazioni" / "Relazione_calcolo.pdf", base / "02_Computi" / "Capitolato.docx")
    cad(base / "03_Disegni")
    ifc(base / "06_BIM")
    images(base / "05_Foto")
    archives(base / "07_Varie")
    cloud_only(base / "07_Varie")
    for f in a.dwg:
        local_copy(f, base / "03_Disegni")
    for f in a.ifc:
        local_copy(f, base / "06_BIM")
    # copie per i tempi: stessi tipi con testo diverso
    srcs = [p for p in base.rglob("*") if p.is_file() and "NUVOLA" not in p.name and p.stat().st_size < 1024 * 1024]
    for i in range(a.copie):
        s = srcs[i % len(srcs)]
        lot = d / f"Archivio_commesse/Lotto_{i // 250:02d}"
        lot.mkdir(parents=True, exist_ok=True)
        if s.suffix.lower() == ".txt":
            (lot / f"{s.stem}_{i}.txt").write_text(" ".join(random.choices(WORDS, k=200)) + f" codice COPIA-{i}", encoding="utf-8")
        else:
            shutil.copy(s, lot / f"{s.stem}_{i}{s.suffix}")
    n = sum(1 for p in d.rglob("*") if p.is_file())
    print(f"Creati {n} file in {d}")


if __name__ == "__main__":
    sys.exit(main())
