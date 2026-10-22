"""Collaudo del recupero dei file danneggiati e sconosciuti (mbvisual/recover.py).

    python tools/make_damaged_samples.py [--dwg disegno.dwg]
    python tests/test_recover.py

Per ogni file di tests/samples/danneggiati (attese in ATTESI.json): controllo veloce, recupero, e verifica che la copia
recuperata si apra davvero con le librerie dei formati (PyMuPDF, python-docx, openpyxl, python-pptx, zipfile, Pillow,
sqlite3, email, FFmpeg). Poi riconoscimento dei file sconosciuti, carving, testo leggibile e la regola "mai sopra
l'originale". Gli originali non devono cambiare (impronta SHA-256 prima e dopo).
"""
from __future__ import annotations

import email
import email.policy
import hashlib
import json
import sqlite3
import subprocess
import sys
import warnings
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from mbvisual import recover  # noqa: E402
from mbvisual.paths import find_ffmpeg  # noqa: E402

D = ROOT / "tests" / "samples" / "danneggiati"
ok = bad = 0


def check(cond, what):
    global ok, bad
    if cond:
        ok += 1
        print("  ok  ", what)
    else:
        bad += 1
        print("  NO  ", what)


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def opens(copy: str, exp: dict) -> tuple[bool, str]:
    p = Path(copy)
    ext = p.suffix.lower().lstrip(".")
    if ext == "pdf":
        import pymupdf
        d = pymupdf.open(p)
        n = d.page_count
        d.close()
        return n >= exp.get("min_pagine", 1), f"{n} pagine"
    if ext == "docx":
        import docx
        doc = docx.Document(str(p))
        txt = "\n".join(x.text for x in doc.paragraphs)
        return exp.get("testo", "") in txt, f"{len(doc.paragraphs)} paragrafi"
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p)
        vals = [c.value for row in wb["Computo"].iter_rows() for c in row]
        return exp.get("celle", "") in vals, f"{len(vals)} celle"
    if ext == "pptx":
        import pptx
        pr = pptx.Presentation(str(p))
        return len(pr.slides) > 0, f"{len(pr.slides)} diapositive"
    if ext in ("odt", "zip"):
        with zipfile.ZipFile(p) as z:
            bad_entry = z.testzip()
            n = len(z.namelist())
        if ext == "odt":
            from lxml import etree
            with zipfile.ZipFile(p) as z:
                etree.fromstring(z.read("content.xml"))
                first = z.namelist()[0]
            return bad_entry is None and first == "mimetype", f"{n} voci, mimetype in testa"
        return bad_entry is None and n >= exp.get("min_voci", 1), f"{n} voci"
    if ext == "png":
        from PIL import Image
        with Image.open(p) as im:
            im.load()
            return im.size[0] == exp.get("larghezza", im.size[0]), f"{im.size}"
    if ext in ("sqlite",):
        con = sqlite3.connect(p)
        n = sum(con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0] for (t,) in con.execute("SELECT name FROM sqlite_master WHERE type='table'"))
        ic = con.execute("PRAGMA integrity_check").fetchone()[0]
        con.close()
        return n >= exp.get("min_righe", 1) and ic == "ok", f"{n} righe, integrita' {ic}"
    if ext == "eml":
        m = email.message_from_bytes(p.read_bytes(), policy=email.policy.default)
        body = m.get_body(("plain", "html"))
        txt = body.get_content() if body else ""
        atts = list(m.iter_attachments())
        return exp.get("testo", "") in txt and len(atts) >= exp.get("allegati", 0), f"{len(atts)} allegati"
    if ext in ("mkv", "mp4", "mp3"):
        r = subprocess.run([find_ffmpeg(), "-hide_banner", "-v", "error", "-i", str(p), "-f", "null", "-"], capture_output=True)
        return r.returncode == 0, "decodifica completa" if r.returncode == 0 else r.stderr.decode()[-200:]
    if ext == "dxf":
        import ezdxf
        doc = ezdxf.readfile(p)
        return len(doc.modelspace()) > 0, f"{len(doc.modelspace())} entita'"
    return True, "(formato non verificato)"


def main():
    att = json.loads((D / "ATTESI.json").read_text(encoding="utf-8"))
    before = {n: sha(D / n) for n in att if (D / n).exists()}
    for name, exp in att.items():
        p = D / name
        if not p.exists():
            continue
        print(name)
        if "famiglia" in exp:
            c = recover.check(str(p))
            if "danneggiato" in exp:
                check(c["danneggiato"] == exp["danneggiato"], f"controllo veloce: danneggiato={c['danneggiato']} ({c.get('codice')})")
            r = recover.recover(str(p))
            check(r["famiglia"] == exp["famiglia"], f"famiglia {r['famiglia']}")
            if exp.get("recuperabile") is False:
                check(r["copia"] is None and r["esito"] == "nulla", "non recuperabile, spiegato: " + (r["messaggi"] or [""])[0][:60])
            elif exp["famiglia"] == "dwg":
                if exp.get("anteprima"):
                    check(bool(r.get("anteprima")) and Path(r["anteprima"]).exists(), "anteprima incorporata estratta")
                if exp.get("alternative"):
                    tipi = sorted(a["tipo"] for a in r["alternative"])
                    check(len(r["alternative"]) >= exp["alternative"] and Path(r["alternative"][0]["path"]).read_bytes()[:3] == b"AC1",
                          f"copie di sicurezza proposte: {tipi}")
            else:
                check(r["copia"] is not None, f"copia recuperata ({r['esito']})")
                if r["copia"]:
                    good, info = opens(r["copia"]["path"], exp)
                    check(good, f"la copia si apre: {info}")
                if exp.get("allegati"):
                    check(len(r["elementi"]) >= exp["allegati"], f"allegati estratti: {len(r['elementi'])}")
        if "identifica" in exp:
            g = recover.identify(str(p))["guesses"]
            check(bool(g) and g[0]["ext"] == exp["identifica"], f"riconosciuto come {g[0]['ext'] if g else '-'}")
            if g:
                cp = recover.as_ext(str(p), g[0]["ext"])["path"]
                check(Path(cp).suffix == "." + exp["identifica"] and sha(Path(cp)) == before[name], "copia con l'estensione giusta")
        if "incorporati" in exp:
            c = recover.carve(str(p))
            exts = [x["ext"] for x in c["elementi"]]
            check(all(e in exts for e in exp["incorporati"]), f"file incorporati: {exts}")
            for x in c["elementi"]:
                if x["ext"] in ("jpg", "png"):
                    from PIL import Image
                    with Image.open(x["path"]) as im:
                        im.load()
            t = recover.extract_text(str(p))
            check(exp["testo"] in Path(t["path"]).read_text(encoding="utf-8"), "testo leggibile estratto")
    # mai sopra l'originale
    print("regole")
    src = D / "relazione_troncata.pdf"
    r = recover.recover(str(src))
    try:
        recover.save_copy(r["copia"]["path"], str(src), str(src))
        check(False, "salvataggio sopra l'originale rifiutato")
    except RuntimeError:
        check(True, "salvataggio sopra l'originale rifiutato")
    try:
        recover.save_copy(str(D / "xref_rotta.pdf"), str(D.parent / "copia_non_ammessa.pdf"), str(src))
        check(False, "si salvano solo i file del recupero")
    except RuntimeError:
        check(True, "si salvano solo i file del recupero")
    after = {n: sha(D / n) for n in before}
    check(after == before, "originali invariati (SHA-256)")
    print(f"\n{ok} prove riuscite, {bad} fallite")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
