"""Collaudo delle pratiche edilizie (mbvisual/edilizia.py, modelli in ui/templates/edilizia).

Uso: python tests/test_edilizia.py   (le visure di prova si rifanno con python tools/make_edilizia_samples.py)
Controlla: lettura dei modelli, condizioni delle caselle, campi compilati dall'anagrafica e riletti dal PDF salvato,
nessun resto della sintassi dei modelli nel testo, DOCX, generazione nella cartella di commessa, fascicolo con indice,
segnalibri e collegamenti, lettura delle visure catastali di prova.
"""
import json
import os
import shutil
import sys
import tempfile
import time
import warnings

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
S = os.path.join(ROOT, "tests", "samples", "edilizia")
warnings.simplefilter("ignore")

import pymupdf  # noqa: E402

from mbvisual import edilizia as E  # noqa: E402

ok = fail = 0


def check(name, cond, info=""):
    global ok, fail
    if cond:
        ok += 1
        print("OK  ", name, info)
    else:
        fail += 1
        print("ERR ", name, info)


def run(name, fn):
    try:
        fn()
    except Exception as e:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        check(name, False, repr(e))


A = json.load(open(os.path.join(S, "esempio.mbcomm"), encoding="utf-8"))
TMP = tempfile.mkdtemp(prefix="mbv-edilizia-")


def fields(path):
    """Campi del PDF salvato: nome -> valore (testo) o True/False (caselle)."""
    out = {}
    with pymupdf.open(path) as d:
        for page in d:
            for w in page.widgets():
                if w.field_type == pymupdf.PDF_WIDGET_TYPE_CHECKBOX:
                    out[w.field_name] = w.field_value not in (False, "Off", "", None)
                else:
                    out[w.field_name] = w.field_value
    return out


def text(path):
    with pymupdf.open(path) as d:
        return "\n".join(p.get_text() for p in d)


def t_cond():
    sc = E.contesto(A)
    check("cond: valore", E._cond("imm.catasto=fabbricati", sc))
    check("cond: elenco di valori", E._cond("int.tipologia=restauro|manutenzione_straordinaria", sc))
    check("cond: negazione", E._cond("!int.titolarita_esclusiva", sc))
    check("cond: diverso", E._cond("int.stato_lavori!=realizzati", sc))
    check("cond: e logico", E._cond("int.sicurezza_81&int.notifica_99", sc) and not E._cond("int.sicurezza_81&int.cantiere_200", sc))
    check("cond: o tra percorsi", E._cond("tt.pdc|imprese", sc) and not E._cond("tt.pdc|tt.scia", sc))
    check("cond: valore vuoto ammesso", E._cond("int.altri_atti=nessuno|", sc) and E._cond("int.cambio_uso=no|", sc))
    check("cond: sempre", E._cond("1", sc))
    check("cond: data italiana", sc["int"]["inizio_lavori"] == "16/11/2026", sc["int"]["inizio_lavori"])
    check("contesto: cognome e nome", sc["tit"]["cognome_nome"] == "Bianchi Giulia")
    check("contesto: altri titolari", [x["cognome"] for x in sc["titolari_altri"]] == ["Verdi"])
    check("contesto: procuratore = progettista se non indicato", sc["pr"]["cognome"] == "Neri")
    check("contesto: nome tipologia", "manutenzione straordinaria" in sc["int"]["tipologia_nome"])


def t_parse():
    names = sorted(f for f in os.listdir(E.TPL_DIR) if f.endswith(".mbm"))
    for n in names:
        meta, blocks = E.parse(E._load_tpl(n))
        check(f"modello {n}", len(blocks) > 0, f"{len(blocks)} blocchi")
    for m in E.MODULI:
        meta, _ = E.parse(E._load_tpl(m[1]))
        check(f"meta {m[0]}", bool(meta.get("titolo") and meta.get("fonte") and meta.get("gruppo")), meta.get("gruppo", ""))
    # sintassi: tabella con "|" dentro le condizioni, campo testo {=x}, <br>
    meta, b = E.parse("@imposta doc PROVA\n| [?a=1|2] uno<br>{x.y:20} | due {=doc} |\n")
    tab = [x for x in b if x[0] == "table"][0]
    check("tabella: celle con | nelle condizioni", len(tab[2][0][1]) == 2)
    toks = E._subst(tab[2][0][1][1], {"doc": "PROVA"})
    check("{=percorso} diventa testo", any(t[0] == "w" and "PROVA" in t[1] for t in toks))


def t_pdf_cila():
    out = os.path.join(TMP, "cila.pdf")
    r = E.render_pdf("cila.mbm", A, out)
    f = fields(out)
    check("CILA: pagine", r["pagine"] >= 6, str(r["pagine"]))
    check("CILA: comune", f.get("com_nome") == "Bergamo")
    check("CILA: titolare", f.get("tit_cognome_nome") == "Bianchi Giulia" and f.get("tit_cf") == "BNCGLI80A41A794K")
    check("CILA: nato il", f.get("tit_nato_il") == "01/01/1980")
    check("CILA: dati catastali", f.get("imm_foglio") == "41" and f.get("imm_particella") == "1234" and f.get("imm_subalterno") == "7")
    check("CILA: titolarita' non esclusiva (a.2)", f.get("a_2") is True and f.get("a_1") is False)
    check("CILA: niente parti comuni (b.1)", f.get("b_1") is True and f.get("b_2") is False)
    check("CILA: inizio lavori (d.1.1)", f.get("d_1_1") is True and f.get("int_inizio_lavori") == "16/11/2026")
    check("CILA: tipologia (1.1)", f.get("1_1") is True and f.get("1_2") is False)
    check("CILA: sicurezza, notifica (i.2, i.3.2)", f.get("i_2") is True and f.get("i_3_2") is True and f.get("i_3_1") is False)
    check("CILA: zona sismica 3 (X.6)", f.get("X_6") is True and f.get("X_7") is False)
    check("CILA: impresa indicata (h.1)", f.get("h_1") is True)
    check("CILA: progettista e albo", f.get("prog_cognome_nome") == "Neri Luca" and f.get("prog_albo_n") == "3456")
    check("CILA: descrizione su piu' righe", (f.get("int_descrizione") or "").startswith("Manutenzione straordinaria dell'appartamento"))
    tx = text(out)
    check("CILA: testo del modulo", "COMUNICAZIONE INIZIO LAVORI ASSEVERATA" in tx and "TOLLERANZA DI CUI ALL" in tx)
    check("CILA: niente spazi non separabili nel testo", "\xa0" not in tx and "\xad" not in tx)
    vuoto = E.render_pdf("cila.mbm", {}, os.path.join(TMP, "cila_vuota.pdf"))
    check("CILA: anagrafica vuota", vuoto["pagine"] >= 6 and not any(x[2] for x in vuoto["campi"] if x[1] not in ("luogo_data",)))


def t_pdf_altri():
    for mid, tpl, kind, _ in E.MODULI:
        if kind != "pdf":
            continue
        out = os.path.join(TMP, mid + ".pdf")
        t0 = time.time()
        r = E.render_pdf(tpl, A, out)
        tx = text(out)
        leak = [s for s in ("{", "}", "[?", "@include", "@se", "<br>", "**") if s in tx]
        names = [x[0] for x in r["campi"]] + [x[0] for x in r["caselle"]]
        check(f"PDF {mid}", r["pagine"] >= 1 and not leak and len(names) == len(set(names)),
              f"{r['pagine']} pag., {len(r['campi'])} campi ({sum(1 for x in r['campi'] if x[2])} compilati), {len(r['caselle'])} caselle, {time.time() - t0:.1f} s"
              + (f", sintassi nel testo: {leak}" if leak else ""))
        with pymupdf.open(out) as d:
            check(f"PDF {mid}: modulo compilabile", d.is_form_pdf and d.is_form_pdf >= len(names) * 0.99, str(d.is_form_pdf))


def t_pdf_contenuti():
    f = fields(os.path.join(TMP, "scia.pdf"))
    check("SCIA: tipologia (c.1)", f.get("c_1") is True and f.get("c_3") is False)
    check("SCIA: SCIA semplice (b.1)", f.get("b_1") is True and f.get("b_2") is False)
    check("SCIA: stato legittimo senza difformita' (f.6.1)", f.get("f_6_1") is True and f.get("f_6_2") is False)
    f = fields(os.path.join(TMP, "scia_alt.pdf"))
    check("SCIA alternativa: c.1 sempre, nessun mutamento d'uso (c.1.1)", f.get("c_1") is True and f.get("c_1_1") is True)
    a = json.loads(json.dumps(A))
    a["intervento"]["superbonus"] = "energetico"
    a["immobile"]["ante67"] = True
    out = os.path.join(TMP, "cilas2.pdf")
    E.render_pdf("cilas.mbm", a, out)
    f = fields(out)
    check("CILAS: efficientamento energetico (1.1)", f.get("1_1") is True and f.get("1_2") is False)
    check("CILAS: costruito prima del 1967 (f.1)", f.get("f_1") is True)
    f = fields(os.path.join(TMP, "soggetti.pdf"))
    vals = set(v for v in f.values() if isinstance(v, str))
    check("Soggetti coinvolti: secondo titolare", "Verdi Paolo" in vals and "VRDPLA78C12A794X" in vals)
    check("Soggetti coinvolti: impresa e regolarita' contributiva", "Edilizia Orobica S.r.l." in vals and "987654" in vals)
    check("Soggetti coinvolti: tecnici", "Neri Luca" in vals and "Gialli Marta" in vals)
    f = fields(os.path.join(TMP, "procura.pdf"))
    vals = set(v for v in f.values() if isinstance(v, str))
    check("Procura: titolari e procuratore", {"Bianchi Giulia", "Verdi Paolo", "Neri Luca"} <= vals)
    f = fields(os.path.join(TMP, "fine.pdf"))
    check("Fine lavori: data e titolo", f.get("int_fine_lavori") == "31/03/2027" and f.get("int_titolo_tipo") == "CILA")
    # titoli edilizi nel quadro f) della SCIA
    a = json.loads(json.dumps(A))
    a["immobile"]["titoli"] = [{"tipo": "pdc", "numero": "123/1998", "data": "1998-05-04"},
                               {"tipo": "altro", "numero": "7", "data": "2001-01-02", "descrizione": "agibilita'"}]
    out = os.path.join(TMP, "scia_titoli.pdf")
    E.render_pdf("scia.mbm", a, out)
    f = fields(out)
    check("SCIA: titolo precedente (f.1.2)", f.get("f_1_2") is True and f.get("tt_pdc_numero") == "123/1998" and f.get("tt_pdc_data") == "04/05/1998")
    check("SCIA: titolo 'altro' (f.1.10)", f.get("f_1_10") is True and f.get("titoli_altri_0_descrizione") == "agibilita'")


def t_docx():
    from docx import Document
    for mid, tpl, kind, _ in E.MODULI:
        if kind != "docx":
            continue
        out = os.path.join(TMP, mid + ".docx")
        r = E.render_docx(tpl, A, out)
        d = Document(out)
        tx = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
        leak = [s for s in ("{", "[?", "@include", "@se ", "**") if s in tx]
        check(f"DOCX {mid}", "Bianchi Giulia" in tx and "41" in tx and not leak, f"{len(r['compilati'])} compilati, {len(r['vuoti'])} da completare" + (f", sintassi: {leak}" if leak else ""))


def t_genera_fascicolo():
    d = os.path.join(TMP, "commessa")
    os.makedirs(d, exist_ok=True)
    mb = os.path.join(d, "Prova.mbcomm")
    shutil.copy(os.path.join(S, "esempio.mbcomm"), mb)
    t0 = time.time()
    r = E.genera(mb, [m[0] for m in E.MODULI])
    check("genera: tutti i moduli", len(r["generati"]) == len(E.MODULI) and not r["errori"], f"{time.time() - t0:.1f} s, cartella {os.path.basename(r['cartella'])}")
    a = E.leggi(mb)
    check("genera: elenco nell'anagrafica", len(a["generati"]) == len(E.MODULI))
    r2 = E.genera(mb, ["cila"])
    check("genera: senza sovrascrivere", r2["generati"][0]["path"].endswith("CILA (2).pdf"), os.path.basename(r2["generati"][0]["path"]))
    os.remove(r2["generati"][0]["path"])
    lst = E.handle("file", {"cartella": r["cartella"]}, None)["file"]
    order = [x["nome"] for x in lst]
    check("cartella: moduli nell'ordine dell'elenco", order[0].endswith("CILA.pdf") and order[1].endswith("CILA-Superbonus.pdf"), ", ".join(order[:3]))
    pdfs = [x["path"] for x in lst if x["nome"].endswith(".pdf")]
    out = os.path.join(d, "Fascicolo.pdf")
    t0 = time.time()
    E.fascicolo(pdfs, out, "Fascicolo di prova", a)
    with pymupdf.open(out) as doc:
        pages = sum(pymupdf.open(p).page_count for p in pdfs)
        toc = doc.get_toc()
        links = doc[0].get_links()
        check("fascicolo: pagine = frontespizio + documenti", doc.page_count == pages + 1, f"{doc.page_count} pagine, {time.time() - t0:.1f} s")
        check("fascicolo: segnalibri", sum(1 for x in toc if x[0] == 1) == len(pdfs) + 1, f"{len(toc)} voci")
        check("fascicolo: indice con collegamenti", len(links) == len(pdfs) and links[0]["page"] == 1, f"{len(links)} collegamenti")
        check("fascicolo: campi appiattiti", not doc.is_form_pdf)
        tx = doc[0].get_text()
        check("fascicolo: frontespizio", "Fascicolo di prova" in tx and "Indice" in tx and "Bianchi Giulia" in tx)
    f2 = E.fascicolo(pdfs[:2] + [x["path"] for x in lst if x["nome"].endswith(".docx")][:1], os.path.join(d, "F2.pdf"), "", a, appiattisci=False)
    check("fascicolo: con DOCX e campi mantenuti", f2["pagine"] > 2 and pymupdf.open(f2["path"]).is_form_pdf)


def t_visura():
    v = E.visura(os.path.join(S, "visura_fabbricati.pdf"))
    imm = v["immobile"]
    check("visura: comune", imm.get("comune") == "Bergamo" and imm.get("catasto") == "fabbricati")
    check("visura: foglio, particella, sub", (imm.get("foglio"), imm.get("particella"), imm.get("subalterno")) == ("41", "1234", "7"))
    check("visura: classamento", (imm.get("categoria"), imm.get("classe"), imm.get("consistenza"), imm.get("rendita")) == ("A/2", "3", "5,5 vani", "795,34"),
          f"{imm.get('categoria')} {imm.get('classe')} {imm.get('consistenza')} {imm.get('rendita')}")
    check("visura: indirizzo", imm.get("via") == "Via Giuseppe Garibaldi" and imm.get("civico") == "12" and imm.get("piano") == "2" and imm.get("interno") == "5")
    it = v["intestati"]
    check("visura: intestati", [(x["cognome"], x["nome"], x["cf"]) for x in it] == [("Bianchi", "Giulia", "BNCGLI80A41A794K"), ("Verdi", "Paolo", "VRDPLA78C12A794X")])
    check("visura: nascita e diritto", it[0]["nato_a"] == "Bergamo" and it[0]["nato_il"] == "1980-01-01" and it[0]["titolo"] == "comproprietario" and it[0]["quota"] == "1/2")
    v = E.visura(os.path.join(S, "visura_terreni.pdf"))
    check("visura terreni", v["immobile"].get("catasto") == "terreni" and v["immobile"].get("foglio") == "5" and "rendita" not in v["immobile"])
    check("visura terreni: societa'", v["intestati"][0]["ditta"]["ragione"] == "IMMOBILIARE OROBICA S.R.L." and v["intestati"][0]["cf"] == "01234567890")
    check("codice fiscale: cognome composto", E._split_name("DE LUCA MARIA ANNA", "DLCMNN80A41A794K") == ("De Luca", "Maria Anna"))
    try:
        blank = os.path.join(TMP, "vuota.pdf")
        d = pymupdf.open()
        d.new_page()
        d.save(blank)
        E.visura(blank)
        check("visura: scansione senza testo rifiutata", False)
    except RuntimeError:
        check("visura: scansione senza testo rifiutata", True)


def t_anagrafica():
    a = E.normalizza({"committenti": [{"cognome": "Rossi"}], "immobile": {"foglio": "3"}})
    check("normalizza: campi mancanti aggiunti", a["committenti"][0]["ditta"]["ragione"] == "" and a["immobile"]["catasto"] == "fabbricati"
          and a["intervento"]["stato_lavori"] == "da_iniziare")
    p = os.path.join(TMP, "x.mbcomm")
    E.scrivi(p, {"commessa": {"nome": "Città è"}})
    check("scrivi/leggi: accenti e struttura", E.leggi(p)["commessa"]["nome"] == "Città è" and open(p, "rb").read().count(b"\r") == 0)
    check("regioni", len(E.REGIONI) == 22 and E.REGIONI[0][0] == "")


for name, fn in [("condizioni", t_cond), ("modelli", t_parse), ("CILA", t_pdf_cila), ("PDF", t_pdf_altri), ("contenuti", t_pdf_contenuti),
                 ("DOCX", t_docx), ("generazione e fascicolo", t_genera_fascicolo), ("visura", t_visura), ("anagrafica", t_anagrafica)]:
    run(name, fn)

shutil.rmtree(TMP, ignore_errors=True)
print(f"\n{ok}/{ok + fail} prove riuscite")
sys.exit(1 if fail else 0)
