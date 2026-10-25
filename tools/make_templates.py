"""Crea i modelli vuoti usati da "Crea nuovo" (ui/templates/). Uso: python tools/make_templates.py"""
import sqlite3
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ui" / "templates"
OUT.mkdir(parents=True, exist_ok=True)



def _no_author(cp) -> None:
    """Modelli senza autore: i documenti nuovi non devono riportare nomi (ne' della libreria ne' di una ditta)."""
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = ""
    cp.comments = ""

def word():
    from docx import Document
    from docx.shared import Cm, Pt
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)  # A4
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(2))
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_paragraph("")
    _no_author(d.core_properties)
    d.save(OUT / "vuoto.docx")


def excel():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Foglio1"
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.save(OUT / "vuoto.xlsx")


def powerpoint():
    from pptx import Presentation
    from pptx.util import Cm
    p = Presentation()
    p.slide_width, p.slide_height = Cm(33.867), Cm(19.05)  # 16:9
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Titolo della presentazione"
    s.placeholders[1].text = "Sottotitolo"
    _no_author(p.core_properties)
    p.save(OUT / "vuoto.pptx")


def pdf():
    import pymupdf
    for name, (w, h) in {"A4": (595, 842), "A4_orizzontale": (842, 595), "A3": (842, 1191), "A3_orizzontale": (1191, 842)}.items():
        d = pymupdf.open()
        d.new_page(width=w, height=h)
        d.save(OUT / f"vuoto_{name}.pdf")


def dxf():
    import ezdxf
    from ezdxf.enums import TextEntityAlignment
    d = ezdxf.new("R2018", setup=True, units=4)  # millimetri
    d.saveas(OUT / "vuoto.dxf")
    # fogli con squadratura e cartiglio semplice (formati ISO orizzontali, margine 10 mm)
    for name, (w, h) in {"A4": (297, 210), "A3": (420, 297), "A2": (594, 420), "A1": (841, 594)}.items():
        d = ezdxf.new("R2018", setup=True, units=4)
        d.layers.add("SQUADRATURA", color=7)
        d.layers.add("CARTIGLIO", color=7)
        d.layers.add("TESTI", color=7)
        msp = d.modelspace()
        msp.add_lwpolyline([(0, 0), (w, 0), (w, h), (0, h)], close=True, dxfattribs={"layer": "SQUADRATURA", "color": 8})
        msp.add_lwpolyline([(10, 10), (w - 10, 10), (w - 10, h - 10), (10, h - 10)], close=True, dxfattribs={"layer": "SQUADRATURA"})
        x0, y0, cw = w - 10 - 180, 10, 180
        msp.add_lwpolyline([(x0, y0), (x0 + cw, y0), (x0 + cw, y0 + 40), (x0, y0 + 40)], close=True, dxfattribs={"layer": "CARTIGLIO"})
        for y in (y0 + 10, y0 + 20, y0 + 30):
            msp.add_line((x0, y), (x0 + cw, y), dxfattribs={"layer": "CARTIGLIO"})
        msp.add_line((x0 + 120, y0), (x0 + 120, y0 + 30), dxfattribs={"layer": "CARTIGLIO"})
        for i, (label, value) in enumerate([("Titolo", ""), ("Commessa", ""), ("Disegnato", ""), ("Scala", "1:1")][::-1]):
            msp.add_text(label.upper(), height=2, dxfattribs={"layer": "CARTIGLIO"}).set_placement((x0 + 2, y0 + i * 10 + 6.5))
        msp.add_text("DITTA / STUDIO", height=2, dxfattribs={"layer": "CARTIGLIO"}).set_placement((x0 + 122, y0 + 6.5))
        msp.add_text(f"Foglio {name}", height=2, dxfattribs={"layer": "CARTIGLIO"}).set_placement((x0 + 122, y0 + 22))
        d.saveas(OUT / f"foglio_{name}.dxf")


def database():
    p = OUT / "vuoto.sqlite"
    p.unlink(missing_ok=True)
    con = sqlite3.connect(p)
    con.execute("CREATE TABLE elenco (id INTEGER PRIMARY KEY, descrizione TEXT, quantita REAL, note TEXT)")
    con.commit()
    con.close()


if __name__ == "__main__":
    for fn in (word, excel, powerpoint, pdf, dxf, database):
        fn()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:28} {f.stat().st_size:>8} byte")
