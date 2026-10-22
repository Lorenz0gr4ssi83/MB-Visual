"""File danneggiati ad arte per il collaudo del recupero (mbvisual/recover.py) -> tests/samples/danneggiati.

    python tools/make_damaged_samples.py [--dwg disegno.dwg]

Ogni file nasce sano (PyMuPDF, python-docx, openpyxl, Pillow, FFmpeg, sqlite3...) e poi si rovina in un modo noto:
troncato, indice cancellato, byte sovrascritti. `tests/samples/danneggiati/ATTESI.json` dice cosa ci si aspetta di
recuperare (lo usa tests/test_recover.py). Con --dwg si aggiungono un DWG troncato, la sua copia .bak e un .sv$
(il DWG non si mette nel repository: serve un disegno a scelta; senza, si crea un DWG finto con la sola anteprima).
"""
from __future__ import annotations

import argparse
import io
import json
import os
import random
import sqlite3
import struct
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "samples" / "danneggiati"
sys.path.insert(0, str(ROOT))
random.seed(7)


def cut(data: bytes, frac: float) -> bytes:
    return data[:int(len(data) * frac)]


# ---------------------------------------------------------------- PDF
def make_pdf(pages=6) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for i in range(pages):
        p = doc.new_page()
        p.insert_text((72, 90), f"Relazione tecnica - pagina {i + 1}", fontsize=20)
        p.insert_text((72, 130), "Calcolo strutturale del solaio in latero-cemento, verifica a flessione.", fontsize=11)
        p.draw_rect(pymupdf.Rect(72, 160, 400, 300), color=(0.9, 0.45, 0.15), width=2)
    data = doc.tobytes(garbage=0, deflate=False, no_new_id=True)
    doc.close()
    return data


def pdf_samples(att: dict):
    full = make_pdf()
    # 1) troncato: mancano xref, trailer e le ultime pagine
    (OUT / "relazione_troncata.pdf").write_bytes(cut(full, 0.8))
    att["relazione_troncata.pdf"] = {"famiglia": "pdf", "danneggiato": True, "min_pagine": 3}
    # 2) xref con posizioni sbagliate e startxref fuori dal file
    b = bytearray(full)
    i = b.rfind(b"startxref")
    b[i:] = b"startxref\n99999999\n%%EOF\n"
    x = b.find(b"xref\n")
    for k in range(x + 30, min(len(b), x + 400)):
        if 0x30 <= b[k] <= 0x39:
            b[k] = 0x30 + random.randrange(10)
    (OUT / "xref_rotta.pdf").write_bytes(bytes(b))
    att["xref_rotta.pdf"] = {"famiglia": "pdf", "danneggiato": True, "min_pagine": 6}
    # 3) catalogo e albero delle pagine cancellati: restano solo gli oggetti sparsi delle pagine
    import re
    b = bytearray(full)
    for m in re.finditer(rb"\d+ 0 obj\s*<<[^>]*?/Type\s*/(Pages|Catalog)\b", bytes(full)):
        end = full.find(b"endobj", m.start())
        b[m.start():end] = b"%" * (end - m.start())
    i = b.rfind(b"xref")
    del b[i:]
    (OUT / "oggetti_sparsi.pdf").write_bytes(bytes(b))
    att["oggetti_sparsi.pdf"] = {"famiglia": "pdf", "danneggiato": True, "min_pagine": 6}


# ---------------------------------------------------------------- Office / ZIP
def central_dir_offset(data: bytes) -> int:
    i = data.rfind(b"PK\x05\x06")
    return struct.unpack("<I", data[i + 16:i + 20])[0]


def docx_bytes() -> bytes:
    import docx
    from PIL import Image
    d = docx.Document()
    d.add_heading("Verbale di sopralluogo", 1)
    for i in range(40):
        d.add_paragraph(f"Punto {i + 1}: verifica dello stato dei luoghi e delle opere eseguite, con annotazioni del direttore dei lavori.")
    im = Image.new("RGB", (320, 200), (234, 116, 38))
    bio = io.BytesIO()
    im.save(bio, "PNG")
    bio.seek(0)
    d.add_picture(bio)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def entry_span(data: bytes, name: str) -> tuple[int, int]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        zi = z.getinfo(name)
    start = zi.header_offset
    nlen, xlen = struct.unpack("<HH", data[start + 26:start + 30])
    return start, start + 30 + nlen + xlen + zi.compress_size


def office_samples(att: dict):
    d = docx_bytes()
    (OUT / "verbale_senza_indice.docx").write_bytes(d[:central_dir_offset(d)])
    att["verbale_senza_indice.docx"] = {"famiglia": "zip", "danneggiato": True, "apre": "docx", "testo": "Verbale di sopralluogo"}
    # parte principale tagliata a meta': si ricostruisce con lxml (resta l'inizio del verbale)
    # (ordine delle voci di python-docx: document.xml non e' l'ultima, quindi si sposta in coda prima di tagliare)
    src = zipfile.ZipFile(io.BytesIO(d))
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for n in sorted(src.namelist(), key=lambda n: n == "word/document.xml"):
            z.writestr(n, src.read(n))
    d2 = bio.getvalue()
    a, b = entry_span(d2, "word/document.xml")
    (OUT / "verbale_troncato.docx").write_bytes(d2[:a + (b - a) // 2])
    att["verbale_troncato.docx"] = {"famiglia": "zip", "danneggiato": True, "apre": "docx", "testo": "Verbale di sopralluogo"}

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Computo"
    ws.append(["Voce", "Quantita", "Prezzo"])
    for i in range(200):
        ws.append([f"Voce {i + 1}", i * 1.5, 12.3 + i])
    ws2 = wb.create_sheet("Riepilogo")
    ws2["A1"] = "Totale"
    ws2["B1"] = "=SUM(Computo!C2:C201)"
    bio = io.BytesIO()
    wb.save(bio)
    x = bytearray(bio.getvalue())
    a, b = entry_span(bytes(x), "xl/styles.xml")
    for k in range(a + 60, b - 4):  # dati compressi degli stili rovinati: la voce si perde, il resto resta
        x[k] ^= 0x5A
    (OUT / "computo_stili_rovinati.xlsx").write_bytes(bytes(x))
    att["computo_stili_rovinati.xlsx"] = {"famiglia": "zip", "danneggiato": True, "apre": "xlsx", "celle": "Voce 200"}

    p = (ROOT / "tests" / "samples" / "presentazione.pptx").read_bytes()
    (OUT / "presentazione_senza_indice.pptx").write_bytes(p[:central_dir_offset(p)])
    att["presentazione_senza_indice.pptx"] = {"famiglia": "zip", "danneggiato": True, "apre": "pptx"}

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text", compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text>'
                   + "".join(f"<text:p>Paragrafo {i + 1} della relazione OpenDocument.</text:p>" for i in range(30))
                   + "</office:text></office:body></office:document-content>")
        z.writestr("styles.xml", '<?xml version="1.0" encoding="UTF-8"?><office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" office:version="1.2"/>')
        z.writestr("META-INF/manifest.xml", '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                   '<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/>'
                   '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                   '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/></manifest:manifest>')
    o = bio.getvalue()
    a, b = entry_span(o, "META-INF/manifest.xml")
    (OUT / "relazione_troncata.odt").write_bytes(o[:a + 10])  # manifest e indice persi
    att["relazione_troncata.odt"] = {"famiglia": "zip", "danneggiato": True, "apre": "odt"}

    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        for i in range(6):
            z.writestr(f"documenti/nota_{i + 1}.txt", f"Nota {i + 1}\n" + "testo di prova " * 400)
        z.writestr("documenti/ultima.bin", os.urandom(30000))
    o = bio.getvalue()
    (OUT / "archivio_troncato.zip").write_bytes(o[:central_dir_offset(o) - 9000])
    att["archivio_troncato.zip"] = {"famiglia": "zip", "danneggiato": True, "min_voci": 6}


# ---------------------------------------------------------------- immagini
def image_samples(att: dict):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (900, 600))
    dr = ImageDraw.Draw(im)
    for y in range(600):
        dr.line([(0, y), (900, y)], fill=(40 + y // 4, 90, 200 - y // 4))
    dr.rectangle([100, 100, 400, 300], fill=(240, 125, 46))
    dr.ellipse([500, 200, 800, 500], fill=(240, 240, 240))
    bio = io.BytesIO()
    im.save(bio, "JPEG", quality=90)
    (OUT / "foto_troncata.jpg").write_bytes(cut(bio.getvalue(), 0.55))
    att["foto_troncata.jpg"] = {"famiglia": "image", "danneggiato": True, "larghezza": 900}
    bio = io.BytesIO()
    im.save(bio, "PNG")
    (OUT / "schema_troncato.png").write_bytes(cut(bio.getvalue(), 0.6))
    att["schema_troncato.png"] = {"famiglia": "image", "danneggiato": True, "larghezza": 900}
    bio = io.BytesIO()
    im.convert("P").save(bio, "GIF")
    (OUT / "animazione_troncata.gif").write_bytes(cut(bio.getvalue(), 0.7))
    att["animazione_troncata.gif"] = {"famiglia": "image", "danneggiato": True, "larghezza": 900}


# ---------------------------------------------------------------- audio / video
def media_samples(att: dict):
    from mbvisual.paths import find_ffmpeg
    ff = find_ffmpeg()
    if not ff:
        print("FFmpeg non trovato: video saltati")
        return
    tmp = OUT / "_tmp"
    tmp.mkdir(exist_ok=True)
    base = [ff, "-hide_banner", "-loglevel", "error", "-y", "-f", "lavfi", "-i", "testsrc=size=320x240:rate=25:duration=8",
            "-f", "lavfi", "-i", "sine=frequency=440:duration=8", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest"]
    fast, slow, mkv = tmp / "fast.mp4", tmp / "slow.mp4", tmp / "v.mkv"
    subprocess.run([*base, "-movflags", "+faststart", str(fast)], check=True)
    subprocess.run([*base, str(slow)], check=True)
    subprocess.run([*base, str(mkv)], check=True)
    (OUT / "ripresa_troncata.mp4").write_bytes(cut(fast.read_bytes(), 0.6))
    att["ripresa_troncata.mp4"] = {"famiglia": "media", "danneggiato": True, "recuperabile": True}
    (OUT / "ripresa_senza_indice.mp4").write_bytes(cut(slow.read_bytes(), 0.6))
    att["ripresa_senza_indice.mp4"] = {"famiglia": "media", "danneggiato": True, "recuperabile": False}
    b = bytearray(mkv.read_bytes())
    for k in range(len(b) // 3, len(b) // 3 + 4000):
        b[k] = random.randrange(256)
    (OUT / "ripresa_rovinata.mkv").write_bytes(bytes(cut(bytes(b), 0.8)))
    att["ripresa_rovinata.mkv"] = {"famiglia": "media", "danneggiato": False, "recuperabile": True}
    for f in tmp.iterdir():
        f.unlink()
    tmp.rmdir()


# ---------------------------------------------------------------- SQLite
def sqlite_samples(att: dict):
    tmp = OUT / "_db.sqlite"
    tmp.unlink(missing_ok=True)
    con = sqlite3.connect(tmp)
    con.execute("PRAGMA page_size=1024")
    con.execute("CREATE TABLE commesse (id INTEGER PRIMARY KEY, codice TEXT, cliente TEXT, importo REAL, note TEXT)")
    con.execute("CREATE TABLE ore (id INTEGER PRIMARY KEY, commessa INTEGER, tecnico TEXT, ore REAL)")
    con.execute("CREATE INDEX ix_ore ON ore(commessa)")
    con.executemany("INSERT INTO commesse VALUES (?,?,?,?,?)",
                    [(i, f"C{i:04d}", f"Cliente {i % 37}", i * 101.5, "nota " * (i % 9)) for i in range(1, 1501)])
    con.executemany("INSERT INTO ore VALUES (?,?,?,?)", [(i, i % 1500 + 1, f"Tecnico {i % 11}", (i % 8) + 0.5) for i in range(1, 4001)])
    con.commit()
    con.close()
    data = tmp.read_bytes()
    tmp.unlink()
    (OUT / "commesse_troncato.sqlite").write_bytes(data[:int(len(data) * 0.7) // 1024 * 1024 + 300])
    att["commesse_troncato.sqlite"] = {"famiglia": "sqlite", "danneggiato": True, "min_righe": 800}
    b = bytearray(data)
    npages = len(b) // 1024
    for pg in random.sample(range(5, npages - 2), 6):  # pagine nel mezzo sovrascritte
        b[pg * 1024:(pg + 1) * 1024] = os.urandom(1024)
    (OUT / "commesse_pagine_rovinate.sqlite").write_bytes(bytes(b))
    att["commesse_pagine_rovinate.sqlite"] = {"famiglia": "sqlite", "danneggiato": True, "min_righe": 4000}


# ---------------------------------------------------------------- email
def email_samples(att: dict):
    import base64
    pdf = make_pdf(2)
    b64 = base64.encodebytes(pdf).decode()
    lines = b64.splitlines()
    lines[len(lines) // 2] = "@@@ riga rovinata ###"
    eml = ("From: Ufficio tecnico <ufficio@example.com>\r\nTo: cantiere@example.com\r\nSubject: Verbale e tavole\r\n"
           "Date: Mon, 20 Oct 2026 09:30:00 +0200\r\nMIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=\"XYZ\"\r\n\r\n"
           "--XYZ\r\nContent-Type: text/plain; charset=utf-8\r\n\r\nIn allegato il verbale del sopralluogo di oggi.\r\n"
           "--XYZ\r\nContent-Type: application/pdf; name=\"verbale.pdf\"\r\nContent-Transfer-Encoding: base64\r\n"
           "Content-Disposition: attachment; filename=\"verbale.pdf\"\r\n\r\n" + "\r\n".join(lines[: int(len(lines) * 0.9)]))
    (OUT / "verbale_rovinato.eml").write_bytes(eml.encode())
    att["verbale_rovinato.eml"] = {"famiglia": "email", "danneggiato": False, "allegati": 1, "testo": "verbale del sopralluogo"}
    msg = make_msg("Riunione di coordinamento", "Direzione lavori", "Ordine del giorno: avanzamento lavori e sicurezza in cantiere.",
                   [("programma.txt", b"Programma lavori settimana 43\n" * 50), ("foto.png", png_small())])
    (OUT / "riunione.msg").write_bytes(msg)
    att["riunione.msg"] = {"famiglia": "email", "danneggiato": False, "allegati": 2, "testo": "Ordine del giorno"}
    b = bytearray(msg)
    b[0x30:0x34] = b"\xff\xff\xff\x7f"  # primo settore della directory fuori dal file
    (OUT / "riunione_rovinata.msg").write_bytes(bytes(b))
    att["riunione_rovinata.msg"] = {"famiglia": "email", "danneggiato": True, "allegati": 1, "testo": "Ordine del giorno"}


def png_small() -> bytes:
    from PIL import Image
    bio = io.BytesIO()
    Image.new("RGB", (64, 48), (30, 140, 90)).save(bio, "PNG")
    return bio.getvalue()


def make_msg(subject: str, sender: str, body: str, atts: list[tuple[str, bytes]]) -> bytes:
    """Messaggio Outlook minimo ([MS-OXMSG]) in un file composto ([MS-CFB]) scritto qui: solo per le prove."""
    def u16(s):
        return s.encode("utf-16-le")
    entries = [("Root Entry", 5, None, [])]  # nome, tipo (5 radice, 1 storage, 2 stream), dati, figli
    root_children = entries[0][3]

    def stream(name, data, parent):
        entries.append((name, 2, data, None))
        parent.append(len(entries) - 1)

    props = bytearray(32)
    stream("__properties_version1.0", bytes(props), root_children)
    stream("__substg1.0_0037001F", u16(subject), root_children)
    stream("__substg1.0_0C1A001F", u16(sender), root_children)
    stream("__substg1.0_1000001F", u16(body), root_children)
    stream("__substg1.0_001A001F", u16("IPM.Note"), root_children)
    named = ("__nameid_version1.0", 1, None, [])
    entries.append(named)
    root_children.append(len(entries) - 1)
    for sid in ("00020102", "00030102", "00040102"):
        stream("__substg1.0_" + sid, b"", named[3])
    for i, (name, data) in enumerate(atts):
        st = (f"__attach_version1.0_#{i:08X}", 1, None, [])
        entries.append(st)
        root_children.append(len(entries) - 1)
        kids = st[3]
        stream("__properties_version1.0", bytes(8 + 8 * 2), kids)
        stream("__substg1.0_3707001F", u16(name), kids)
        stream("__substg1.0_37010102", data, kids)
    SEC, MINI = 512, 64
    mini_data = bytearray()
    mini_fat: list[int] = []
    big: list[tuple[int, bytes]] = []  # (indice voce, dati)
    starts = {}
    for idx, (name, typ, data, _) in enumerate(entries):
        if typ != 2:
            continue
        if not data:
            starts[idx] = 0xFFFFFFFE
        elif len(data) < 4096:
            n = max(1, -(-len(data) // MINI))
            first = len(mini_fat)
            mini_fat += [first + k + 1 for k in range(n - 1)] + [0xFFFFFFFE]
            mini_data += data + b"\x00" * (n * MINI - len(data))
            starts[idx] = first
        else:
            big.append((idx, data))
    # settori: [directory][ministream][minifat][big streams][fat]
    sectors: list[bytes] = []
    fat: list[int] = []

    def alloc(data: bytes) -> int:
        n = max(1, -(-len(data) // SEC))
        first = len(sectors)
        for k in range(n):
            sectors.append(data[k * SEC:(k + 1) * SEC].ljust(SEC, b"\x00"))
            fat.append(first + k + 1 if k < n - 1 else 0xFFFFFFFE)
        return first

    ndir = -(-len(entries) // 4)
    dir_first = alloc(b"\x00" * (ndir * SEC))
    mini_first = alloc(bytes(mini_data)) if mini_data else 0xFFFFFFFE
    mf = b"".join(struct.pack("<I", v) for v in mini_fat)
    minifat_first = alloc(mf)
    for idx, data in big:
        starts[idx] = alloc(data)
    nfat = 1
    while (len(sectors) + nfat) * 4 > nfat * SEC:
        nfat += 1
    fat_first = len(sectors)
    for k in range(nfat):
        fat.append(0xFFFFFFFD)
        sectors.append(b"")
    fat_bytes = b"".join(struct.pack("<I", v) for v in fat).ljust(nfat * SEC, b"\xff")
    for k in range(nfat):
        sectors[fat_first + k] = fat_bytes[k * SEC:(k + 1) * SEC]

    # directory: figli di ogni storage in un albero (catena a destra, ordinati come vuole CFB: lunghezza poi maiuscole)
    def key(i):
        n = entries[i][0]
        return (len(n), n.upper())
    child_of, right = {}, {}
    for idx, (_, typ, _, kids) in enumerate(entries):
        if kids:
            ks = sorted(kids, key=key)
            child_of[idx] = ks[0]
            for a, b2 in zip(ks, ks[1:]):
                right[a] = b2
    dirb = bytearray()
    for idx, (name, typ, data, kids) in enumerate(entries):
        nm = u16(name)[:62] + b"\x00\x00"
        e = bytearray(128)
        e[:len(nm)] = nm
        struct.pack_into("<H", e, 64, len(nm))
        e[66] = typ
        e[67] = 1
        struct.pack_into("<III", e, 68, 0xFFFFFFFF, right.get(idx, 0xFFFFFFFF), child_of.get(idx, 0xFFFFFFFF))
        if typ == 5:
            struct.pack_into("<IQ", e, 116, mini_first, len(mini_data))
        elif typ == 2:
            struct.pack_into("<IQ", e, 116, starts[idx], len(data))
        dirb += e
    dirb = dirb.ljust(ndir * SEC, b"\x00")
    for k in range(ndir):
        sectors[dir_first + k] = bytes(dirb[k * SEC:(k + 1) * SEC])
        if k >= len(entries) // 4 + (1 if len(entries) % 4 else 0):
            pass
    hdr = bytearray(512)
    hdr[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", hdr, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIII", hdr, 40, 0, nfat, dir_first, 0, 4096, minifat_first, max(1, -(-len(mf) // SEC)), 0xFFFFFFFE)
    struct.pack_into("<I", hdr, 72, 0)
    difat = [fat_first + k for k in range(nfat)] + [0xFFFFFFFF] * (109 - nfat)
    struct.pack_into("<109I", hdr, 76, *difat)
    # voci libere della directory
    return bytes(hdr) + b"".join(sectors)


# ---------------------------------------------------------------- DWG
DWG_SENTINEL = bytes.fromhex("1F256D07D43628289D57CA3F9D44102B")


def fake_dwg_with_preview() -> bytes:
    """Intestazione DWG R2018 con l'anteprima BMP all'indirizzo di 0x0D; il resto e' rovinato."""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (160, 120), (0, 0, 0))
    dr = ImageDraw.Draw(im)
    dr.rectangle([20, 20, 140, 100], outline=(255, 255, 255))
    dr.line([(20, 20), (140, 100)], fill=(255, 255, 0))
    bio = io.BytesIO()
    im.save(bio, "BMP")
    dib = bio.getvalue()[14:]
    head = bytearray(b"AC1032" + b"\x00" * 0x7A)
    addr = 0x100
    struct.pack_into("<I", head, 0x0D, addr)
    body = bytearray(os.urandom(addr - len(head)))
    img = DWG_SENTINEL + struct.pack("<I", 1 + 9 * 2 + len(dib)) + bytes([2])
    start = addr + len(img) + 18
    img += bytes([1]) + struct.pack("<II", start, 0) + bytes([2]) + struct.pack("<II", start, len(dib))
    return bytes(head) + bytes(body) + img + dib + os.urandom(5000)


def dwg_samples(att: dict, dwg: str | None):
    for f in OUT.glob("disegno*"):
        f.unlink()
    if dwg:
        data = Path(dwg).read_bytes()
        (OUT / "disegno_troncato.dwg").write_bytes(cut(data, 0.85))
        (OUT / "disegno_troncato.bak").write_bytes(data)
        (OUT / "disegno_troncato_1_1_4821.sv$").write_bytes(data)
        att["disegno_troncato.dwg"] = {"famiglia": "dwg", "alternative": 2}
    (OUT / "disegno_rovinato.dwg").write_bytes(fake_dwg_with_preview())
    att["disegno_rovinato.dwg"] = {"famiglia": "dwg", "anteprima": True}


# ---------------------------------------------------------------- file sconosciuti
def unknown_samples(att: dict):
    pdf = make_pdf(1)
    (OUT / "documento_senza_estensione").write_bytes(pdf)
    att["documento_senza_estensione"] = {"identifica": "pdf"}
    (OUT / "foto_rinominata.xyz").write_bytes(png_small())
    att["foto_rinominata.xyz"] = {"identifica": "png"}
    (OUT / "elenco.dat").write_bytes("codice;descrizione;quantita\n1;Trave HEB 200;12\n2;Pilastro;4\n3;Solaio;120\n".encode())
    att["elenco.dat"] = {"identifica": "csv"}
    from PIL import Image
    bio = io.BytesIO()
    Image.new("RGB", (200, 150), (234, 116, 38)).save(bio, "JPEG")
    blob = (os.urandom(3000) + b"intestazione proprietaria MBV-TEST " + os.urandom(2000) + bio.getvalue() + os.urandom(1500)
            + png_small() + os.urandom(800) + make_pdf(1) + os.urandom(700))
    (OUT / "contenitore_ignoto.bin").write_bytes(blob)
    att["contenitore_ignoto.bin"] = {"incorporati": ["jpg", "png", "pdf"], "testo": "intestazione proprietaria"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dwg", help="DWG sano da cui ricavare disegno_troncato.dwg con .bak e .sv$ (non va nel repository)")
    a = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    att: dict = {}
    pdf_samples(att)
    office_samples(att)
    image_samples(att)
    media_samples(att)
    sqlite_samples(att)
    email_samples(att)
    dwg_samples(att, a.dwg)
    unknown_samples(att)
    (OUT / "ATTESI.json").write_bytes(json.dumps(att, indent=1, ensure_ascii=False).encode() + b"\n")
    print(f"{len(att)} file in {OUT}")


if __name__ == "__main__":
    main()
