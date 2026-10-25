"""Crea file di esempio in tests/samples per collaudare ogni categoria di MB Desk."""
import datetime
import io
import json
import sqlite3
import struct
import subprocess
import wave
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "samples"
OUT.mkdir(parents=True, exist_ok=True)


def pdf_bytes():
    pdf = ROOT / "ui" / "vendor" / "pdfjs" / "web" / "compressed.tracemonkey-pldi-09.pdf"
    return pdf.read_bytes()


def make_p7m():
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs7
    from cryptography.x509.oid import NameOID

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([
        x509.NameAttribute(NameOID.COUNTRY_NAME, "IT"),
        x509.NameAttribute(NameOID.GIVEN_NAME, "MARIO"),
        x509.NameAttribute(NameOID.SURNAME, "ROSSI"),
        x509.NameAttribute(NameOID.SERIAL_NUMBER, "TINIT-RSSMRA80A01H501U"),
        x509.NameAttribute(NameOID.COMMON_NAME, "ROSSI MARIO"),
    ])
    issuer = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "CA DI PROVA MB DESK"),
                        x509.NameAttribute(NameOID.ORGANIZATION_NAME, "Certificatore di prova")])
    now = datetime.datetime.now(datetime.timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(issuer).public_key(key.public_key())
            .serial_number(x509.random_serial_number()).not_valid_before(now - datetime.timedelta(days=1))
            .not_valid_after(now + datetime.timedelta(days=365))
            .add_extension(x509.KeyUsage(digital_signature=False, content_commitment=True, key_encipherment=False,
                                         data_encipherment=False, key_agreement=False, key_cert_sign=False,
                                         crl_sign=False, encipher_only=False, decipher_only=False), critical=True)
            .sign(key, hashes.SHA256()))
    data = pdf_bytes()
    der = (pkcs7.PKCS7SignatureBuilder().set_data(data).add_signer(cert, key, hashes.SHA256())
           .sign(serialization.Encoding.DER, [pkcs7.PKCS7Options.Binary]))
    (OUT / "contratto_firmato.pdf.p7m").write_bytes(der)


def make_office():
    from docx import Document
    d = Document()
    d.add_heading("Relazione tecnica di prova", 0)
    d.add_paragraph("Documento di esempio generato per il collaudo di MB Desk.")
    p = d.add_paragraph()
    p.add_run("Testo in grassetto ").bold = True
    p.add_run("e testo in corsivo.").italic = True
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"R{r + 1}C{c + 1}"
    d.add_paragraph("Primo punto", style="List Bullet")
    d.add_paragraph("Secondo punto", style="List Bullet")
    d.save(OUT / "relazione.docx")

    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Presentazione di prova"
    s.placeholders[1].text = "MB Desk - collaudo PowerPoint"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Punti principali"
    tf = s2.placeholders[1].text_frame
    tf.text = "Apertura di qualsiasi file"
    tf.add_paragraph().text = "Modifica dei testi"
    box = s2.shapes.add_textbox(Inches(1), Inches(5.5), Inches(6), Inches(1))
    box.text_frame.text = "Casella di testo libera"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    prs.save(OUT / "presentazione.pptx")

    import xlsxwriter
    wb = xlsxwriter.Workbook(str(OUT / "computo.xlsx"))
    ws = wb.add_worksheet("Computo")
    ws.write_row(0, 0, ["Voce", "Quantita", "Prezzo", "Totale"])
    for i, (v, q, pz) in enumerate([("Calcestruzzo m3", 12, 110), ("Acciaio kg", 850, 1.4), ("Casseri m2", 64, 22)], 1):
        ws.write_row(i, 0, [v, q, pz])
        ws.write_formula(i, 3, f"=B{i + 1}*C{i + 1}")
    ws.write_formula(4, 3, "=SUM(D2:D4)")
    wb.add_worksheet("Note").write(0, 0, "Secondo foglio")
    wb.close()
    (OUT / "dati.csv").write_text("nome;valore\nalfa;1\nbeta;2\n", encoding="utf-8")


def make_cad():
    dxf = """0
SECTION
2
ENTITIES
0
LINE
8
MURI
10
0
20
0
11
100
21
0
0
LINE
8
MURI
10
100
20
0
11
100
21
60
0
LINE
8
MURI
10
100
20
60
11
0
21
60
0
LINE
8
MURI
10
0
20
60
11
0
21
0
0
CIRCLE
8
FORI
10
50
20
30
40
15
0
TEXT
8
TESTI
10
5
20
65
40
5
1
PIANTA DI PROVA
0
ENDSEC
0
EOF
"""
    (OUT / "pianta.dxf").write_text(dxf, encoding="ascii")
    # cubo STL ASCII
    v = [(0, 0, 0), (10, 0, 0), (10, 10, 0), (0, 10, 0), (0, 0, 10), (10, 0, 10), (10, 10, 10), (0, 10, 10)]
    f = [(0, 2, 1), (0, 3, 2), (4, 5, 6), (4, 6, 7), (0, 1, 5), (0, 5, 4), (1, 2, 6), (1, 6, 5), (2, 3, 7), (2, 7, 6), (3, 0, 4), (3, 4, 7)]
    lines = ["solid cubo"]
    for a, b, c in f:
        lines += ["facet normal 0 0 0", " outer loop", *[f"  vertex {x} {y} {z}" for x, y, z in (v[a], v[b], v[c])], " endloop", "endfacet"]
    lines.append("endsolid cubo")
    (OUT / "cubo.stl").write_text("\n".join(lines), encoding="ascii")
    (OUT / "cubo.obj").write_text("\n".join([f"v {x} {y} {z}" for x, y, z in v] + [f"f {a + 1} {b + 1} {c + 1}" for a, b, c in f]), encoding="ascii")


def make_media():
    with wave.open(str(OUT / "tono.wav"), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(22050)
        import math
        w.writeframes(b"".join(struct.pack("<h", int(12000 * math.sin(2 * math.pi * 440 * i / 22050) * (1 - i / 66150))) for i in range(66150)))
    import imageio_ffmpeg
    ff = imageio_ffmpeg.get_ffmpeg_exe()
    subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=4:size=640x360:rate=25",
                    "-f", "lavfi", "-i", "sine=frequency=330:duration=4", "-c:v", "mpeg4", "-c:a", "mp2", str(OUT / "video_prova.avi")], check=True)
    subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=duration=3:size=640x360:rate=25",
                    "-c:v", "libx264", "-pix_fmt", "yuv420p", str(OUT / "video_prova.mp4")], check=True)


def make_misc():
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (640, 400), (230, 236, 240))
    d = ImageDraw.Draw(im)
    d.rectangle([40, 40, 600, 360], outline=(47, 102, 144), width=6)
    d.ellipse([220, 120, 420, 320], fill=(138, 140, 137))
    im.save(OUT / "immagine.png")
    im.save(OUT / "immagine.tif")
    im.save(OUT / "immagine.tga")
    with zipfile.ZipFile(OUT / "archivio.zip", "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("leggimi.txt", "Archivio di prova MB Desk\n")
        z.writestr("cartella/dati.json", json.dumps({"a": 1, "b": [1, 2, 3]}, indent=2))
        z.write(OUT / "immagine.png", "cartella/immagine.png")
    eml = ("From: Ufficio Tecnico <tecnico@example.com>\r\nTo: Cliente <cliente@example.com>\r\nSubject: Invio elaborati\r\n"
           "Date: Tue, 20 Oct 2026 10:00:00 +0200\r\nMIME-Version: 1.0\r\nContent-Type: multipart/mixed; boundary=XX\r\n\r\n"
           "--XX\r\nContent-Type: text/html; charset=utf-8\r\n\r\n<p>Buongiorno,<br>in allegato gli <b>elaborati</b>.</p>\r\n"
           "--XX\r\nContent-Type: text/plain; name=nota.txt\r\nContent-Disposition: attachment; filename=nota.txt\r\n\r\nNota allegata\r\n--XX--\r\n")
    (OUT / "messaggio.eml").write_bytes(eml.encode())
    db = OUT / "archivio_commesse.sqlite"
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE commesse (id INTEGER PRIMARY KEY, codice TEXT, cliente TEXT, importo REAL)")
    con.executemany("INSERT INTO commesse (codice, cliente, importo) VALUES (?,?,?)",
                    [("C-2026-001", "Comune di Prova", 125000), ("C-2026-002", "Alfa Spa", 48000.5)])
    con.commit()
    con.close()
    (OUT / "note.md").write_text("# Note di progetto\n\n- **Fase 1**: rilievo\n- Fase 2: progetto\n\n| Voce | Stato |\n|---|---|\n| Rilievo | fatto |\n", encoding="utf-8")
    (OUT / "pagina.html").write_text("<!doctype html><html><body><h1>Pagina di prova</h1><p>Testo</p></body></html>", encoding="utf-8")
    (OUT / "disegno.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="120"><rect x="10" y="10" width="180" height="100" fill="#8a8c89"/><text x="30" y="70" fill="#fff" font-size="24">MB Desk</text></svg>', encoding="utf-8")
    (OUT / "sconosciuto.xyz").write_bytes(bytes(range(256)) * 4)


if __name__ == "__main__":
    make_p7m()
    make_office()
    make_cad()
    make_media()
    make_misc()
    print("\n".join(sorted(p.name for p in OUT.iterdir())))
